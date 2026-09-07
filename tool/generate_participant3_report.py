"""Создаёт персональный отчёт по практике для третьего участника."""

from pathlib import Path

from docx import Document
from docx.enum.table import (
    WD_CELL_VERTICAL_ALIGNMENT,
    WD_ROW_HEIGHT_RULE,
    WD_TABLE_ALIGNMENT,
)
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Отчет_участника_3_исправленный.docx"


def set_font(run, size=14, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="808080", size="8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(3)
    section.right_margin = Cm(1.5)
    section.header_distance = Cm(1)
    section.footer_distance = Cm(1)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(14)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(0)

    for level, size in ((1, 16), (2, 14), (3, 14)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT
        )
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def add_field(paragraph, instruction, fallback=""):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph._p.append(begin)
    paragraph._p.append(code)
    if fallback:
        run = paragraph.add_run(fallback)
        set_font(run, 12)
    paragraph._p.append(separate)
    paragraph._p.append(end)


def add_footer(doc):
    p = doc.sections[0].footer.paragraphs[0]
    p.clear()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    add_field(p, "PAGE", "1")


def centered(doc, value, size=14, bold=False, before=0, after=0):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    set_font(p.add_run(value), size, bold)
    return p


def add_title_page(doc):
    centered(doc, "МИНИСТЕРСТВО НАУКИ, ВЫСШЕГО ОБРАЗОВАНИЯ", 12, True)
    centered(doc, "И ИННОВАЦИЙ КЫРГЫЗСКОЙ РЕСПУБЛИКИ", 12, True)
    centered(doc, "КЫРГЫЗСКИЙ ГОСУДАРСТВЕННЫЙ ТЕХНИЧЕСКИЙ", 12, True, before=8)
    centered(doc, "УНИВЕРСИТЕТ им. И. РАЗЗАКОВА", 12, True)
    centered(doc, "ИНСТИТУТ ИНФОРМАЦИОННЫХ ТЕХНОЛОГИЙ", 12, True, before=8)
    centered(doc, "Кафедра «Программное обеспечение компьютерных систем»", 12)
    centered(doc, "ОТЧЁТ", 20, True, before=70)
    centered(doc, "ПО ПРОИЗВОДСТВЕННОЙ ПРАКТИКЕ", 16, True, after=18)
    centered(doc, "на тему:", 14)
    centered(
        doc,
        "«РАЗРАБОТКА И ТЕСТИРОВАНИЕ МУЛЬТИМЕДИЙНЫХ И РЕЛИГИОЗНЫХ "
        "МОДУЛЕЙ ОБРАЗОВАТЕЛЬНОГО ПРИЛОЖЕНИЯ IRFAN ACADEMY»",
        15,
        True,
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.left_indent = Cm(8.5)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(55)
    for line in (
        "Выполнил: студент группы ПИ ан-2-23",
        "[ФИО ТРЕТЬЕГО УЧАСТНИКА]",
        "_____________________________",
        "Место практики: [ОРГАНИЗАЦИЯ]",
        "Руководитель от организации:",
        "[ФИО, ДОЛЖНОСТЬ] ____________",
        "Руководитель от кафедры:",
        "[ФИО, ДОЛЖНОСТЬ] ____________",
    ):
        set_font(p.add_run(line + "\n"))
    centered(doc, "Бишкек — 2026", 14, before=65)
    doc.add_page_break()


def add_individual_assignment(doc):
    doc.add_heading("ИНДИВИДУАЛЬНОЕ ЗАДАНИЕ", level=1)
    p = doc.add_paragraph()
    p.add_run("Студент: ").bold = True
    p.add_run("[ФИО ТРЕТЬЕГО УЧАСТНИКА], группа ПИ ан-2-23.")
    p = doc.add_paragraph()
    p.add_run("Роль в проекте: ").bold = True
    p.add_run("разработчик мультимедийных и религиозных модулей, ответственный за проверку реализованных им функций.")
    p = doc.add_paragraph()
    p.add_run("Тема индивидуальной работы: ").bold = True
    p.add_run("разработка и тестирование загрузки и воспроизведения видео, геолокации, времени намаза, направления Кыблы и локальных напоминаний.")

    tasks = [
        ("1", "Исследовать варианты внешнего хранения видеоуроков."),
        ("2", "Реализовать выбор и загрузку видео с отображением прогресса."),
        ("3", "Обеспечить обработку крупных файлов, ошибок, отмены и повторных попыток."),
        ("4", "Подключить настоящее воспроизведение сетевого видео."),
        ("5", "Получать координаты пользователя и рассчитывать время намаза."),
        ("6", "Рассчитать направление Кыблы и связать его с данными компаса."),
        ("7", "Реализовать локальные напоминания о предстоящих намазах."),
        ("8", "Проверить разработанные модули в Android- и Web-версиях."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "№"
    table.rows[0].cells[1].text = "Содержание задания"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "D9EAF7")
    for number, task in tasks:
        cells = table.add_row().cells
        cells[0].text = number
        cells[1].text = task
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p2 in cell.paragraphs:
                p2.paragraph_format.first_line_indent = Cm(0)
                p2.paragraph_format.line_spacing = 1.0
                for run in p2.runs:
                    set_font(run, 12, row_index == 0)
    doc.add_paragraph()
    for label in (
        "Период практики: [ДАТА НАЧАЛА] — [ДАТА ОКОНЧАНИЯ]",
        "Дата выдачи задания: [ДАТА]",
        "Подпись студента: ____________________",
        "Подпись руководителя: ________________",
    ):
        p = doc.add_paragraph(label)
        p.paragraph_format.first_line_indent = Cm(0)
    doc.add_page_break()


def add_toc(doc):
    doc.add_heading("ОГЛАВЛЕНИЕ", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    add_field(p, 'TOC \\o "1-3" \\h \\z \\u', "Оглавление обновляется в Microsoft Word")
    note = doc.add_paragraph("После окончательного заполнения отчёта следует выделить оглавление и нажать F9 для обновления номеров страниц.")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.first_line_indent = Cm(0)
    for run in note.runs:
        set_font(run, 11, italic=True)
    doc.add_page_break()


def text(doc, value):
    doc.add_paragraph(value)


def bullets(doc, values):
    for value in values:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.first_line_indent = Cm(0)
        p.add_run(value)


def code(doc, value):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F2F2")
    set_cell_border(cell, "B7B7B7", "6")
    p = cell.paragraphs[0]
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(value), 10, name="Courier New")


def screenshot(doc, number, description, height=6.0):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    row = table.rows[0]
    row.height = Cm(height)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = row.cells[0]
    cell.width = Cm(15.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, "F2F2F2")
    set_cell_border(cell)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    set_font(p.add_run(f"МЕСТО ДЛЯ СКРИНШОТА {number}\n\n{description}"), 12, True)
    caption = doc.add_paragraph(f"Рисунок {number} — {description}")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.first_line_indent = Cm(0)
    for run in caption.runs:
        set_font(run, 12, italic=True)


def add_technology_table(doc):
    rows = (
        ("Flutter/Dart", "реализация модулей для Android и Web"),
        ("Cloudinary", "загрузка и выдача сетевого URL видео"),
        ("image_picker", "выбор видеофайла с устройства"),
        ("http", "потоковая и частичная отправка файла"),
        ("video_player", "реальное воспроизведение видео"),
        ("geolocator", "получение координат пользователя"),
        ("adhan", "расчёт намазов и направления Кыблы"),
        ("flutter_compass", "получение направления устройства"),
        ("flutter_local_notifications", "локальные напоминания о намазе"),
    )
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.rows[0].cells[0].text = "Инструмент"
    table.rows[0].cells[1].text = "Назначение в индивидуальной работе"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "D9EAF7")
    for tool, purpose in rows:
        cells = table.add_row().cells
        cells[0].text = tool
        cells[1].text = purpose
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_font(run, 11, row_index == 0)


def add_test_table(doc):
    rows = (
        ("Т-01", "Выбор корректного видео", "Файл выбран, загрузка запускается", "Пройден"),
        ("Т-02", "Отмена выбора", "Форма остаётся в исходном состоянии", "Пройден"),
        ("Т-03", "Прогресс загрузки", "Процент изменяется до завершения", "Пройден"),
        ("Т-04", "Файл больше 95 МБ", "Отправляется частями по 8 МБ", "Пройден"),
        ("Т-05", "Временный сетевой сбой", "Часть отправляется повторно", "Пройден"),
        ("Т-06", "Отмена загрузки", "Процесс прекращается без публикации", "Пройден"),
        ("Т-07", "Открытие URL видео", "Плеер инициализируется", "Пройден"),
        ("Т-08", "Пауза и перемотка", "Позиция меняется корректно", "Пройден"),
        ("Т-09", "Запрет геолокации", "Используются координаты Бишкека", "Пройден"),
        ("Т-10", "Расчёт после Иша", "Следующим выбран Фаджр нового дня", "Пройден"),
        ("Т-11", "Кыбла и компас", "Указатель реагирует на поворот", "Пройден"),
        ("Т-12", "Запуск Web-версии", "zonedSchedule не вызывается", "Пройден"),
        ("Т-13", "Android-напоминание", "Уведомление планируется", "Пройден"),
    )
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, value in zip(table.rows[0].cells, ("Код", "Проверка", "Ожидаемый результат", "Статус")):
        cell.text = value
        set_cell_shading(cell, "D9EAF7")
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.line_spacing = 1.0
                for run in p.runs:
                    set_font(run, 9, row_index == 0)


def build_report():
    doc = Document()
    configure_document(doc)
    add_title_page(doc)
    add_individual_assignment(doc)
    add_toc(doc)

    doc.add_heading("ВВЕДЕНИЕ", level=1)
    text(doc, "В период производственной практики выполнялась разработка мультимедийных и религиозных модулей образовательного приложения Irfan Academy. Практическая часть была сосредоточена на функциях, использующих внешнее видеохранилище и возможности устройства пользователя.")
    text(doc, "Индивидуальная работа была связана с функциями, зависящими от внешних сервисов и возможностей устройства: загрузкой и воспроизведением видеоуроков, определением координат, вычислением времени намаза и направления Кыблы, планированием локальных напоминаний. Отдельной частью работы стала проверка устойчивости этих функций в Android- и Web-версиях приложения.")
    text(doc, "Цель работы — реализовать и проверить мультимедийные и религиозные модули Irfan Academy таким образом, чтобы пользователь мог загружать и просматривать продолжительные видео, получать актуальное расписание намазов и использовать направление Кыблы без нарушения работы существующей части приложения.")

    doc.add_heading("1 ПОСТАНОВКА ИНДИВИДУАЛЬНОЙ ЗАДАЧИ", level=1)
    doc.add_heading("1.1 Границы ответственности", level=2)
    text(doc, "В зоне ответственности находились файлы и сервисы, непосредственно связанные с мультимедиа, геолокацией, религиозными вычислениями и локальными уведомлениями. Реализованные компоненты подключались к существующим точкам приложения и проверялись как самостоятельные функциональные модули.")
    bullets(doc, ["сервис загрузки видео в Cloudinary;", "передача результата загрузки в существующую систему сохранения метаданных;", "виджет сетевого видеоплеера;", "сервис получения координат и расчёта времени намаза;", "виджет направления Кыблы;", "сервис локальных напоминаний о намазе;", "проверка перечисленных функций на Android и в браузере."])
    doc.add_heading("1.2 Требования к результату", level=2)
    text(doc, "Видео должно выбираться с устройства без ручного ввода ссылки. Во время передачи пользователь должен видеть ход операции и иметь возможность отменить её. Продолжительность, размер, формат, URL и идентификатор должны поступать из ответа видеосервиса. Для длинных видео требовалась частичная отправка с повторными попытками.")
    text(doc, "Расписание намазов должно учитывать координаты пользователя и ханафитский мазхаб. При недоступной геолокации приложение должно продолжать работу с безопасным значением по умолчанию — координатами Бишкека. В Web-версии запрещалось вызывать неподдерживаемое планирование локальных уведомлений.")
    doc.add_heading("1.3 Использованные инструменты", level=2)
    add_technology_table(doc)
    screenshot(doc, 1, "Файлы индивидуальных модулей в структуре проекта", 5.5)

    doc.add_heading("2 РАЗРАБОТКА МОДУЛЯ ВИДЕОУРОКОВ", level=1)
    doc.add_heading("2.1 Анализ варианта хранения видео", level=2)
    text(doc, "На начальном этапе были рассмотрены Firebase Storage и Cloudflare R2. Для R2 был подготовлен вариант безопасной серверной схемы с временными URL, при которой ключи доступа не передаются во Flutter-приложение. Практическое развёртывание потребовало подключения платных серверных возможностей. Поэтому для демонстрационной версии был выбран Cloudinary с отдельным unsigned upload preset. Исследование R2 осталось подготовленным вариантом развития, а рабочая реализация проекта использует Cloudinary.")
    text(doc, "Видеофайл не помещается в документ базы данных. Сервис возвращает HTTPS-адрес и идентификатор public_id, а приложение передаёт эти метаданные в существующую логику публикации урока. Благодаря этому большие бинарные данные не проходят через Firestore.")
    screenshot(doc, 2, "Настройка Cloudinary upload preset для демонстрационной загрузки")
    doc.add_heading("2.2 Выбор файла и автоматический запуск", level=2)
    text(doc, "Для выбора видео используется XFile из пакета image_picker. После подтверждения выбора приложение получает длину файла и сразу начинает отправку. Поле ручного ввода URL и поле продолжительности в секундах были удалены из сценария загрузки: пользователь вводит только учебные сведения, а технические параметры определяются автоматически.")
    code(doc, "final total = await file.length();\nif (total <= 0) {\n  throw const FormatException('Выбран пустой видеофайл.');\n}")
    screenshot(doc, 3, "Выбор видеофайла с устройства")
    doc.add_heading("2.3 Потоковая и частичная загрузка", level=2)
    text(doc, "Файлы размером до 95 МБ отправляются потоковым MultipartRequest без предварительного чтения всего содержимого в память. Количество отправленных байтов используется для вычисления процента. Это уменьшает расход памяти и позволяет показывать фактический ход операции.")
    text(doc, "Для более крупных файлов применяется разбиение на части по 8 МБ. Каждая часть отправляется с общим X-Unique-Upload-Id и диапазоном Content-Range. При временной ошибке выполняется до четырёх попыток с задержками 1, 2 и 4 секунды. Перед каждой частью проверяется признак отмены. Такой подход повышает устойчивость при нестабильном соединении, хотя полное возобновление после закрытия приложения в текущей версии не реализовано.")
    code(doc, "static const chunkSize = 8 * 1024 * 1024;\nstatic const directUploadLimit = 95 * 1024 * 1024;\nstatic const maxAttempts = 4;")
    screenshot(doc, 4, "Индикатор и процент загрузки видео")
    screenshot(doc, 5, "Сообщение об ошибке, отмене или повторной попытке")
    doc.add_heading("2.4 Обработка результата", level=2)
    text(doc, "После завершения Cloudinary возвращает secure_url, public_id, bytes, format и duration. Эти значения преобразуются в CloudinaryUploadResult. Продолжительность округляется до секунд, поэтому администратор больше не рассчитывает её вручную. Если обязательные поля отсутствуют, загрузка считается незавершённой и пользователю выводится понятное сообщение.")
    code(doc, "CloudinaryUploadResult(\n  secureUrl: data['secure_url'],\n  publicId: data['public_id'],\n  fileSize: data['bytes'],\n  durationSeconds: data['duration'].round(),\n)")
    screenshot(doc, 6, "Успешно загруженный видеоурок и автоматически полученные данные")
    doc.add_heading("2.5 Реальное воспроизведение", level=2)
    text(doc, "Имитация воспроизведения таймером была заменена VideoPlayerController.networkUrl. Плеер инициализируется только при наличии корректного HTTPS-адреса. Реализованы запуск, пауза, перемотка, повтор после завершения и изменение скорости. Отдельные состояния используются для инициализации, буферизации и ошибки открытия ресурса.")
    text(doc, "Перед созданием нового контроллера предыдущий экземпляр освобождается. Проверки mounted предотвращают обновление уже закрытого окна. Такой порядок уменьшает вероятность утечки ресурсов и исключений при быстром закрытии плеера.")
    screenshot(doc, 7, "Воспроизведение загруженного видео в сетевом плеере")

    doc.add_heading("3 РАЗРАБОТКА РЕЛИГИОЗНЫХ МОДУЛЕЙ", level=1)
    doc.add_heading("3.1 Получение геолокации", level=2)
    text(doc, "PrayerService проверяет, включена ли геолокация, затем анализирует разрешение и при необходимости запрашивает его. Координаты запрашиваются с высокой точностью и ограничением ожидания пять секунд. Ошибка датчика, запрет разрешения или тайм-аут не останавливают экран: применяется резервное местоположение Бишкек с координатами 42.8746 и 74.5698.")
    text(doc, "Резервный сценарий был добавлен осознанно: пользователь должен видеть расписание даже при первом запуске, отключённом GPS или работе в браузере без разрешения на местоположение.")
    screenshot(doc, 8, "Расписание, рассчитанное по текущей или резервной геолокации")
    doc.add_heading("3.2 Расчёт времени намаза", level=2)
    text(doc, "Для вычислений используется библиотека adhan. В качестве метода выбран Muslim World League, для Асра установлен Madhab.hanafi. Сервис формирует Фаджр, Восход, Зухр, Аср, Магриб и Иша, переводит время в локальный часовой пояс и определяет состояние каждого пункта: завершён, следующий, предстоящий или вспомогательный.")
    text(doc, "Особое внимание уделено переходу суток. Если после Иша в сегодняшнем расписании нет будущего обязательного намаза, сервис отдельно рассчитывает Фаджр следующего дня. Обратный отсчёт поэтому не показывает отрицательное значение и не возвращается к уже прошедшему Фаджру текущего дня.")
    screenshot(doc, 9, "Следующий намаз и обратный отсчёт")
    doc.add_heading("3.3 Направление Кыблы", level=2)
    text(doc, "Абсолютный азимут Кыблы вычисляется классом Qibla на основании тех же координат, что используются для расписания. Flutter Compass предоставляет направление верхней части устройства. Разность этих значений определяет угол поворота указателя. Виджет учитывает отсутствие датчика и не завершает работу с ошибкой, если браузер или устройство не выдаёт heading.")
    screenshot(doc, 10, "Работа указателя направления Кыблы")
    doc.add_heading("3.4 Локальные напоминания", level=2)
    text(doc, "NotificationService инициализирует локальную временную зону и запрашивает разрешения платформы. Перед формированием нового расписания старые напоминания удаляются, после чего создаются уведомления только для будущих обязательных намазов. Сначала запрашивается точное планирование; если платформа его не разрешает, применяется неточный режим.")
    text(doc, "Изначально Web-версия завершалась исключением Unsupported operation: zonedSchedule() is not supported on the web. Причина заключалась в вызове мобильного API в браузере. В initialize и schedulePrayerNotifications добавлена ранняя проверка kIsWeb, поэтому неподдерживаемый код в браузере больше не выполняется.")
    screenshot(doc, 11, "Локальное напоминание о времени намаза на Android")
    screenshot(doc, 12, "Запуск Web-версии без ошибки zonedSchedule")

    doc.add_heading("4 ПРОВЕРКА РАЗРАБОТАННЫХ МОДУЛЕЙ", level=1)
    doc.add_heading("4.1 Организация проверки", level=2)
    text(doc, "Проверка проводилась по сценариям, относящимся только к индивидуальной части работы. Для сетевых функций проверялись успешный ответ, отказ сервиса, обрыв соединения и отмена пользователем. Для платформенных функций проверялись предоставленное и запрещённое разрешение, наличие и отсутствие датчика, а также различия Android и Web.")
    text(doc, "После изменений запускались статический анализ Flutter и существующий набор автоматизированных регрессионных тестов проекта. На момент подготовки отчёта набор содержит два теста, оба выполняются успешно. Они подтверждают отсутствие базовой регрессии, но не заменяют ручные интеграционные сценарии видео, компаса и системных уведомлений, поскольку эти функции используют реальные платформенные API.")
    add_test_table(doc)
    doc.add_heading("4.2 Исправленные дефекты", level=2)
    bullets(doc, ["устранена CORS-ошибка прежнего способа загрузки видео из браузера;", "удалён ручной ввод URL и продолжительности видео;", "снижено потребление памяти за счёт потоковой передачи;", "добавлена частичная отправка больших файлов и повтор временно неудачных частей;", "таймер-имитация заменён реальным сетевым видеоплеером;", "устранён вызов zonedSchedule в Web-версии;", "добавлен переход к Фаджру следующего дня после Иша;", "добавлено резервное местоположение при недоступной геолокации;", "обработано отсутствие данных компаса и ошибки инициализации видео."])
    screenshot(doc, 13, "Успешный результат flutter analyze и flutter test")
    screenshot(doc, 14, "Проверка разработанных модулей в Android- или Web-сборке")

    doc.add_heading("5 РЕЗУЛЬТАТЫ ИНДИВИДУАЛЬНОЙ РАБОТЫ", level=1)
    text(doc, "В результате практики реализован рабочий модуль выбора, передачи и воспроизведения видеоуроков. Загрузка показывает прогресс, поддерживает отмену, повторяет временно неудачные запросы и получает технические характеристики из ответа Cloudinary. Воспроизведение выполняется настоящим сетевым плеером, а не локальным таймером.")
    text(doc, "Реализован сервис расчёта расписания намазов по координатам пользователя с ханафитским Асром, корректным выбором следующего намаза и переходом на новый день. Рассчитанное направление Кыблы связано с показаниями компаса. Для Android создаются локальные напоминания, а Web-версия безопасно исключает неподдерживаемый вызов.")
    text(doc, "Разработанные компоненты объединены в завершённый набор функций: подготовка видео к публикации, получение его технических данных, сетевое воспроизведение, вычисление расписания намазов и Кыблы, а также безопасное планирование напоминаний с учётом платформы запуска.")
    doc.add_heading("5.1 Ограничения текущей реализации", level=2)
    bullets(doc, ["unsigned upload preset подходит для демонстрации, но для закрытого учебного контента потребуется серверная подпись;", "частичная загрузка повторяется в пределах текущего запуска, но не восстанавливается после полного закрытия приложения;", "качество видео не переключается автоматически, поскольку HLS и транскодирование нескольких вариантов не внедрены;", "компас зависит от физического датчика и точности его калибровки;", "интеграционные сценарии платформенных API пока фиксируются ручным протоколом."])
    doc.add_heading("5.2 Направления дальнейшей работы", level=2)
    text(doc, "Следующим этапом для данных модулей целесообразно добавить подписанную загрузку через сервер, сохранение состояния multipart upload, автоматизированные тесты сервисов с подменой HTTP-клиента и геолокации, а также отдельный тестовый контур. Для длинных уроков возможно формирование HLS с несколькими качествами. Эти пункты являются предложениями и не выдаются за выполненную часть практики.")

    doc.add_heading("ЗАКЛЮЧЕНИЕ", level=1)
    text(doc, "В ходе производственной практики была выполнена индивидуальная задача по разработке и проверке мультимедийных и религиозных модулей Irfan Academy. Практическая работа включала исследование внешнего видеохранилища, реализацию устойчивой загрузки и настоящего воспроизведения видео, получение геолокации, расчёт намазов и Кыблы, а также планирование локальных напоминаний.")
    text(doc, "В процессе работы были устранены ошибки, характерные для кроссплатформенного приложения: ограничения CORS, неподдерживаемые браузером уведомления, нестабильная передача больших файлов, отсутствие разрешений и датчиков. Получен опыт работы с потоковыми HTTP-запросами, платформенными API Flutter, обработкой асинхронных ошибок и проверкой поведения программы на разных платформах.")
    text(doc, "Поставленная индивидуальная цель достигнута: разработанные участником функции подключены к проекту и могут быть продемонстрированы независимо от задач, выполненных другими членами команды.")

    doc.add_heading("СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ", level=1)
    sources = (
        "Flutter. Документация по разработке и тестированию приложений. — URL: https://docs.flutter.dev/",
        "Dart. Руководство по языку и асинхронному программированию. — URL: https://dart.dev/guides",
        "Cloudinary. Документация Upload API. — URL: https://cloudinary.com/documentation/image_upload_api_reference",
        "Cloudinary. Загрузка крупных файлов частями. — URL: https://cloudinary.com/documentation/upload_images#chunked_asset_upload",
        "Flutter package: video_player. — URL: https://pub.dev/packages/video_player",
        "Flutter package: image_picker. — URL: https://pub.dev/packages/image_picker",
        "Flutter package: geolocator. — URL: https://pub.dev/packages/geolocator",
        "Flutter package: adhan. — URL: https://pub.dev/packages/adhan",
        "Flutter package: flutter_compass. — URL: https://pub.dev/packages/flutter_compass",
        "Flutter package: flutter_local_notifications. — URL: https://pub.dev/packages/flutter_local_notifications",
    )
    for index, source in enumerate(sources, 1):
        p = doc.add_paragraph(f"{index}. {source}")
        p.paragraph_format.first_line_indent = Cm(0)

    doc.add_heading("ПРИЛОЖЕНИЕ А. ПЕРЕЧЕНЬ МАТЕРИАЛОВ ДЛЯ ЗАЩИТЫ", level=1)
    text(doc, "К отчёту рекомендуется приложить только материалы, подтверждающие индивидуальную работу третьего участника:")
    bullets(doc, ["экран выбора видео и индикатор загрузки;", "ответ Cloudinary без отображения секретных данных;", "работающий видеоплеер;", "расписание намазов и следующий намаз;", "экран направления Кыблы;", "локальное уведомление на Android;", "Web-консоль без ошибки zonedSchedule;", "результат статического анализа и тестов;", "краткий протокол проверок из раздела 4."])
    text(doc, "На скриншотах следует скрыть токены, API-ключи, адреса электронной почты пользователей и другие конфиденциальные данные. Каждое изображение должно иметь номер, подпись и ссылку на него в тексте отчёта.")
    add_footer(doc)
    return doc


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_report()
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()

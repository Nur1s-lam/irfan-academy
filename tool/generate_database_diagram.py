"""Generate the Irfan Academy Firestore database diagram as PNG, SVG and DOCX."""

from __future__ import annotations

from dataclasses import dataclass
from html import escape
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
PNG_PATH = DOCS / "Диаграмма_базы_данных_Irfan_Academy.png"
SVG_PATH = DOCS / "Диаграмма_базы_данных_Irfan_Academy.svg"
DOCX_PATH = DOCS / "Диаграмма_базы_данных_Irfan_Academy.docx"

WIDTH = 3400
HEIGHT = 2300
HEADER_HEIGHT = 82
ROW_HEIGHT = 43


@dataclass(frozen=True)
class Entity:
    key: str
    title: str
    x: int
    y: int
    width: int
    fields: tuple[str, ...]
    color: str
    badge: str

    @property
    def height(self) -> int:
        return HEADER_HEIGHT + ROW_HEIGHT * len(self.fields) + 18

    @property
    def left(self) -> int:
        return self.x

    @property
    def right(self) -> int:
        return self.x + self.width

    @property
    def top(self) -> int:
        return self.y

    @property
    def bottom(self) -> int:
        return self.y + self.height

    @property
    def center_x(self) -> int:
        return self.x + self.width // 2

    @property
    def center_y(self) -> int:
        return self.y + self.height // 2


ENTITIES = {
    "auth": Entity(
        "auth",
        "Firebase Authentication",
        70,
        190,
        650,
        (
            "uid : string [PK]",
            "email : string",
            "displayName : string?",
            "credentials : managed by Firebase",
        ),
        "#455A64",
        "ВНЕШНИЙ СЕРВИС",
    ),
    "users": Entity(
        "users",
        "users/{uid}",
        1040,
        120,
        920,
        (
            "uid : document ID [PK]",
            "name : string",
            "email : string",
            "group : string",
            "role : string",
            "attendance : int",
            "lessonsCount : int",
            "streakDays : int",
            "quranProgress : double",
            "currentSura : string",
            "currentAyah : int",
            "createdAt : timestamp",
        ),
        "#1F4E78",
        "КОЛЛЕКЦИЯ",
    ),
    "cloudinary": Entity(
        "cloudinary",
        "Cloudinary Video Asset",
        2630,
        170,
        700,
        (
            "public_id : string [PK]",
            "secure_url : URL",
            "bytes : int",
            "format : string",
            "duration : number",
        ),
        "#2E6E5B",
        "ВНЕШНЕЕ ХРАНИЛИЩЕ",
    ),
    "announcements": Entity(
        "announcements",
        "announcements/{announcementId}",
        70,
        650,
        690,
        (
            "announcementId : document ID [PK]",
            "title : string",
            "body : string",
            "isPublished : bool",
            "createdAt : timestamp",
        ),
        "#C56A16",
        "КОЛЛЕКЦИЯ",
    ),
    "videos": Entity(
        "videos",
        "videoLessons/{videoId}",
        2470,
        600,
        860,
        (
            "videoId : document ID [PK]",
            "title : string",
            "number : int",
            "totalSeconds : int",
            "duration : string (MM:SS)",
            "videoUrl : URL [REF → secure_url]",
            "videoPath : string [REF → public_id]",
            "fileName : string",
            "fileSize : int",
            "contentType : string",
            "storageProvider : 'cloudinary'",
            "description : string",
            "createdAt : timestamp",
        ),
        "#6A4C93",
        "КОЛЛЕКЦИЯ",
    ),
    "lessons": Entity(
        "lessons",
        "users/{uid}/lessons/{lessonId}",
        70,
        1430,
        650,
        (
            "lessonId : document ID [PK]",
            "day : string",
            "time : string",
            "durationMin : int",
            "type : string",
            "topic : string",
            "teacher : string",
            "isSoon : bool",
            "sortOrder : int",
            "createdAt : timestamp",
        ),
        "#2F75B5",
        "ПОДКОЛЛЕКЦИЯ",
    ),
    "homework": Entity(
        "homework",
        "users/{uid}/homework/{homeworkId}",
        790,
        1430,
        650,
        (
            "homeworkId : document ID [PK]",
            "task : string",
            "subject : string",
            "deadline : string",
            "status : string",
            "isDone : bool",
            "createdAt : timestamp",
        ),
        "#2F75B5",
        "ПОДКОЛЛЕКЦИЯ",
    ),
    "notifications": Entity(
        "notifications",
        "users/{uid}/notifications/{notificationId}",
        1510,
        1430,
        690,
        (
            "notificationId : document ID [PK]",
            "title : string",
            "message : string",
            "type : string",
            "isRead : bool",
            "createdAt : timestamp",
        ),
        "#2F75B5",
        "ПОДКОЛЛЕКЦИЯ",
    ),
    "bookmarks": Entity(
        "bookmarks",
        "users/{uid}/quranBookmarks/{bookmarkId}",
        2270,
        1510,
        820,
        (
            "bookmarkId : document ID [PK]",
            "ayahNumber : int",
            "surahName : string",
            "createdAt : timestamp",
        ),
        "#2F75B5",
        "ПОДКОЛЛЕКЦИЯ",
    ),
}


def rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    return tuple(int(value[i : i + 2], 16) for i in (0, 2, 4))


def load_font(size: int, bold: bool = False, mono: bool = False):
    candidates = []
    if mono:
        candidates.extend((Path("C:/Windows/Fonts/consola.ttf"), Path("C:/Windows/Fonts/cour.ttf")))
    elif bold:
        candidates.extend((Path("C:/Windows/Fonts/arialbd.ttf"), Path("C:/Windows/Fonts/calibrib.ttf")))
    else:
        candidates.extend((Path("C:/Windows/Fonts/arial.ttf"), Path("C:/Windows/Fonts/calibri.ttf")))
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    return ImageFont.load_default()


FONT_TITLE = load_font(54, bold=True)
FONT_SUBTITLE = load_font(28)
FONT_HEADER = load_font(30, bold=True)
FONT_HEADER_MEDIUM = load_font(24, bold=True)
FONT_HEADER_SMALL = load_font(22, bold=True)
FONT_FIELD = load_font(24, mono=True)
FONT_LABEL = load_font(23, bold=True)
FONT_NOTE = load_font(23)
FONT_LEGEND = load_font(24)


def draw_entity(draw: ImageDraw.ImageDraw, entity: Entity):
    shadow = (entity.x + 10, entity.y + 12, entity.right + 10, entity.bottom + 12)
    draw.rounded_rectangle(shadow, radius=20, fill=(0, 0, 0, 32))
    draw.rounded_rectangle(
        (entity.left, entity.top, entity.right, entity.bottom),
        radius=20,
        fill="white",
        outline=rgb(entity.color),
        width=4,
    )
    draw.rounded_rectangle(
        (entity.left, entity.top, entity.right, entity.top + HEADER_HEIGHT),
        radius=20,
        fill=rgb(entity.color),
    )
    draw.rectangle(
        (entity.left, entity.top + HEADER_HEIGHT - 20, entity.right, entity.top + HEADER_HEIGHT),
        fill=rgb(entity.color),
    )
    header_font = (
        FONT_HEADER_SMALL
        if len(entity.title) > 42
        else FONT_HEADER_MEDIUM
        if len(entity.title) > 31
        else FONT_HEADER
    )
    draw.text((entity.left + 22, entity.top + 22), entity.title, font=header_font, fill="white")
    for index, field in enumerate(entity.fields):
        y = entity.top + HEADER_HEIGHT + index * ROW_HEIGHT
        if index % 2 == 0:
            draw.rectangle((entity.left + 3, y, entity.right - 3, y + ROW_HEIGHT), fill="#F5F8FB")
        if index:
            draw.line((entity.left + 3, y, entity.right - 3, y), fill="#D9E1E8", width=1)
        field_fill = "#9C2F2F" if "[PK]" in field else "#263238"
        draw.text((entity.left + 22, y + 8), field, font=FONT_FIELD, fill=field_fill)


def arrow(draw: ImageDraw.ImageDraw, points, color, width=7, dashed=False):
    if dashed:
        for a, b in zip(points, points[1:]):
            x1, y1 = a
            x2, y2 = b
            length = max(abs(x2 - x1), abs(y2 - y1))
            if length == 0:
                continue
            for start in range(0, length, 28):
                end = min(start + 15, length)
                sx = x1 + (x2 - x1) * start / length
                sy = y1 + (y2 - y1) * start / length
                ex = x1 + (x2 - x1) * end / length
                ey = y1 + (y2 - y1) * end / length
                draw.line((sx, sy, ex, ey), fill=color, width=width)
    else:
        draw.line(points, fill=color, width=width, joint="curve")
    (x1, y1), (x2, y2) = points[-2], points[-1]
    if abs(x2 - x1) > abs(y2 - y1):
        direction = 1 if x2 > x1 else -1
        head = [(x2, y2), (x2 - 24 * direction, y2 - 15), (x2 - 24 * direction, y2 + 15)]
    else:
        direction = 1 if y2 > y1 else -1
        head = [(x2, y2), (x2 - 15, y2 - 24 * direction), (x2 + 15, y2 - 24 * direction)]
    draw.polygon(head, fill=color)


def label(draw: ImageDraw.ImageDraw, position, value, color="#263238"):
    x, y = position
    box = draw.textbbox((x, y), value, font=FONT_LABEL)
    draw.rounded_rectangle((box[0] - 10, box[1] - 6, box[2] + 10, box[3] + 6), radius=7, fill="white")
    draw.text((x, y), value, font=FONT_LABEL, fill=color)


def draw_connections(draw: ImageDraw.ImageDraw):
    user = ENTITIES["users"]
    auth = ENTITIES["auth"]
    cloud = ENTITIES["cloudinary"]
    videos = ENTITIES["videos"]
    announcements = ENTITIES["announcements"]
    notifications = ENTITIES["notifications"]

    arrow(draw, [(auth.right, 315), (user.left, 315)], "#526D82")
    label(draw, (780, 270), "uid  1 : 1")

    trunk_y = 1320
    draw.line((user.center_x, user.bottom, user.center_x, trunk_y), fill="#2F75B5", width=8)
    children = [ENTITIES[key] for key in ("lessons", "homework", "notifications", "bookmarks")]
    draw.line((children[0].center_x, trunk_y, children[-1].center_x, trunk_y), fill="#2F75B5", width=8)
    for child in children:
        arrow(draw, [(child.center_x, trunk_y), (child.center_x, child.top)], "#2F75B5")
        label(draw, (child.center_x - 75, trunk_y + 18), "0..*")
    label(draw, (user.center_x + 18, user.bottom + 55), "1 — владелец")

    arrow(draw, [(cloud.center_x, cloud.bottom), (cloud.center_x, 540), (videos.center_x, 540), (videos.center_x, videos.top)], "#2E6E5B", dashed=True)
    label(draw, (2715, 505), "public_id / secure_url")

    event_color = "#D17A22"
    arrow(draw, [(announcements.right, announcements.center_y), (850, announcements.center_y), (850, 1240), (notifications.center_x - 80, 1240), (notifications.center_x - 80, notifications.top)], event_color, dashed=True)
    label(draw, (920, 1195), "публикация → уведомления всем")

    arrow(draw, [(videos.left, videos.center_y), (2320, videos.center_y), (2320, 1265), (notifications.center_x + 80, 1265), (notifications.center_x + 80, notifications.top)], event_color, dashed=True)
    label(draw, (2050, 1217), "новое видео → уведомления")

    homework = ENTITIES["homework"]
    homework_route_y = 1930
    arrow(draw, [(homework.center_x, homework.bottom), (homework.center_x, homework_route_y), (notifications.center_x, homework_route_y), (notifications.center_x, notifications.bottom)], event_color, dashed=True)
    label(draw, (1110, homework_route_y - 44), "задание → notification(type='homework')")

    lessons = ENTITIES["lessons"]
    route_y = 2070
    arrow(draw, [(lessons.center_x, lessons.bottom), (lessons.center_x, route_y), (notifications.center_x, route_y), (notifications.center_x, notifications.bottom)], event_color, dashed=True)
    label(draw, (850, route_y - 45), "назначение урока → notification(type='lesson')")


def create_png():
    image = Image.new("RGBA", (WIDTH, HEIGHT), "#EEF3F7")
    draw = ImageDraw.Draw(image, "RGBA")
    draw.text((70, 35), "ЛОГИЧЕСКАЯ СХЕМА FIRESTORE — IRFAN ACADEMY", font=FONT_TITLE, fill="#17324D")
    draw.text((72, 100), "Коллекции, подколлекции, внешнее хранилище и фактические связи по коду проекта", font=FONT_SUBTITLE, fill="#526D82")
    draw_connections(draw)
    for entity in ENTITIES.values():
        draw_entity(draw, entity)

    legend_y = 2140
    draw.rounded_rectangle((70, legend_y, 3330, 2260), radius=18, fill="white", outline="#B6C5D1", width=3)
    draw.line((100, legend_y + 38, 210, legend_y + 38), fill="#2F75B5", width=8)
    draw.text((230, legend_y + 22), "владение / вложенная подколлекция", font=FONT_LEGEND, fill="#263238")
    arrow(draw, [(800, legend_y + 38), (910, legend_y + 38)], "#D17A22", width=6, dashed=True)
    draw.text((930, legend_y + 22), "логическая связь / создаваемое уведомление", font=FONT_LEGEND, fill="#263238")
    draw.text((1780, legend_y + 22), "[PK] — идентификатор документа; [REF] — логическая ссылка", font=FONT_LEGEND, fill="#263238")
    draw.text((100, legend_y + 72), "Не хранятся в Firestore: видеофайл — Cloudinary; время намаза и Кыбла — вычисляются; Коран — пакет quran; настройки уведомлений — SharedPreferences.", font=FONT_NOTE, fill="#455A64")
    image.convert("RGB").save(PNG_PATH, quality=95, dpi=(220, 220))


def svg_text(x, y, value, size, weight="400", fill="#263238", family="Arial"):
    return f'<text x="{x}" y="{y}" font-family="{family}" font-size="{size}" font-weight="{weight}" fill="{fill}">{escape(value)}</text>'


def svg_entity(entity: Entity):
    parts = [
        f'<rect x="{entity.x + 10}" y="{entity.y + 12}" width="{entity.width}" height="{entity.height}" rx="20" fill="#000" opacity="0.12"/>',
        f'<rect x="{entity.x}" y="{entity.y}" width="{entity.width}" height="{entity.height}" rx="20" fill="#fff" stroke="{entity.color}" stroke-width="4"/>',
        f'<path d="M {entity.x + 20} {entity.y} H {entity.right - 20} Q {entity.right} {entity.y} {entity.right} {entity.y + 20} V {entity.y + HEADER_HEIGHT} H {entity.x} V {entity.y + 20} Q {entity.x} {entity.y} {entity.x + 20} {entity.y} Z" fill="{entity.color}"/>',
        svg_text(
            entity.x + 22,
            entity.y + 52,
            entity.title,
            22 if len(entity.title) > 42 else 24 if len(entity.title) > 31 else 30,
            "700",
            "#fff",
        ),
    ]
    for index, field in enumerate(entity.fields):
        y = entity.y + HEADER_HEIGHT + index * ROW_HEIGHT
        if index % 2 == 0:
            parts.append(f'<rect x="{entity.x + 3}" y="{y}" width="{entity.width - 6}" height="{ROW_HEIGHT}" fill="#F5F8FB"/>')
        if index:
            parts.append(f'<line x1="{entity.x + 3}" y1="{y}" x2="{entity.right - 3}" y2="{y}" stroke="#D9E1E8"/>')
        parts.append(svg_text(entity.x + 22, y + 29, field, 24, "400", "#9C2F2F" if "[PK]" in field else "#263238", "Consolas, monospace"))
    return "\n".join(parts)


def svg_polyline(points, color, dashed=False, marker=True):
    value = " ".join(f"{x},{y}" for x, y in points)
    dash = ' stroke-dasharray="16 13"' if dashed else ""
    end = ' marker-end="url(#arrow)"' if marker else ""
    return f'<polyline points="{value}" fill="none" stroke="{color}" stroke-width="7" stroke-linejoin="round"{dash}{end}/>'


def create_svg():
    u = ENTITIES["users"]
    a = ENTITIES["auth"]
    c = ENTITIES["cloudinary"]
    v = ENTITIES["videos"]
    ann = ENTITIES["announcements"]
    n = ENTITIES["notifications"]
    h = ENTITIES["homework"]
    lesson = ENTITIES["lessons"]
    children = [ENTITIES[key] for key in ("lessons", "homework", "notifications", "bookmarks")]
    parts = [
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{WIDTH}" height="{HEIGHT}" viewBox="0 0 {WIDTH} {HEIGHT}">',
        '<defs><marker id="arrow" viewBox="0 0 10 10" refX="9" refY="5" markerWidth="9" markerHeight="9" orient="auto-start-reverse"><path d="M 0 0 L 10 5 L 0 10 z" fill="context-stroke"/></marker></defs>',
        f'<rect width="{WIDTH}" height="{HEIGHT}" fill="#EEF3F7"/>',
        svg_text(70, 78, "ЛОГИЧЕСКАЯ СХЕМА FIRESTORE — IRFAN ACADEMY", 54, "700", "#17324D"),
        svg_text(72, 126, "Коллекции, подколлекции, внешнее хранилище и фактические связи по коду проекта", 28, "400", "#526D82"),
        svg_polyline([(a.right, 315), (u.left, 315)], "#526D82"),
        svg_text(780, 300, "uid  1 : 1", 23, "700", "#263238"),
        svg_polyline([(u.center_x, u.bottom), (u.center_x, 1320)], "#2F75B5", marker=False),
        svg_polyline([(children[0].center_x, 1320), (children[-1].center_x, 1320)], "#2F75B5", marker=False),
    ]
    for child in children:
        parts.append(svg_polyline([(child.center_x, 1320), (child.center_x, child.top)], "#2F75B5"))
        parts.append(svg_text(child.center_x - 30, 1375, "0..*", 23, "700", "#263238"))
    parts.extend(
        [
            svg_text(u.center_x + 18, u.bottom + 75, "1 — владелец", 23, "700", "#263238"),
            svg_polyline([(c.center_x, c.bottom), (c.center_x, 540), (v.center_x, 540), (v.center_x, v.top)], "#2E6E5B", True),
            svg_text(2715, 530, "public_id / secure_url", 23, "700", "#263238"),
            svg_polyline([(ann.right, ann.center_y), (850, ann.center_y), (850, 1240), (n.center_x - 80, 1240), (n.center_x - 80, n.top)], "#D17A22", True),
            svg_text(920, 1225, "публикация → уведомления всем", 23, "700", "#263238"),
            svg_polyline([(v.left, v.center_y), (2320, v.center_y), (2320, 1265), (n.center_x + 80, 1265), (n.center_x + 80, n.top)], "#D17A22", True),
            svg_text(2050, 1245, "новое видео → уведомления", 23, "700", "#263238"),
            svg_polyline([(h.center_x, h.bottom), (h.center_x, 1930), (n.center_x, 1930), (n.center_x, n.bottom)], "#D17A22", True),
            svg_text(1110, 1915, "задание → notification(type='homework')", 23, "700", "#263238"),
            svg_polyline([(lesson.center_x, lesson.bottom), (lesson.center_x, 2070), (n.center_x, 2070), (n.center_x, n.bottom)], "#D17A22", True),
            svg_text(850, 2055, "назначение урока → notification(type='lesson')", 23, "700", "#263238"),
        ]
    )
    parts.extend(svg_entity(entity) for entity in ENTITIES.values())
    parts.extend(
        [
            '<rect x="70" y="2140" width="3260" height="120" rx="18" fill="#fff" stroke="#B6C5D1" stroke-width="3"/>',
            '<line x1="100" y1="2178" x2="210" y2="2178" stroke="#2F75B5" stroke-width="8"/>',
            svg_text(230, 2187, "владение / вложенная подколлекция", 24),
            svg_polyline([(800, 2178), (910, 2178)], "#D17A22", True),
            svg_text(930, 2187, "логическая связь / создаваемое уведомление", 24),
            svg_text(1780, 2187, "[PK] — идентификатор документа; [REF] — логическая ссылка", 24),
            svg_text(100, 2233, "Не хранятся в Firestore: видеофайл — Cloudinary; время намаза и Кыбла — вычисляются; Коран — пакет quran; настройки уведомлений — SharedPreferences.", 23, "400", "#455A64"),
            "</svg>",
        ]
    )
    SVG_PATH.write_text("\n".join(parts), encoding="utf-8")


def set_run_font(run, size, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold


def create_docx():
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(420)
    section.page_height = Mm(297)
    section.top_margin = Cm(1)
    section.bottom_margin = Cm(1)
    section.left_margin = Cm(1)
    section.right_margin = Cm(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    set_run_font(title.add_run("Диаграмма базы данных Irfan Academy"), 16, True)

    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.space_after = Pt(3)
    picture.add_run().add_picture(str(PNG_PATH), width=Cm(38.5))

    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(0)
    set_run_font(caption.add_run("Рисунок 1 — Логическая схема коллекций и связей Cloud Firestore"), 12)
    doc.save(DOCX_PATH)


def main():
    DOCS.mkdir(parents=True, exist_ok=True)
    create_png()
    create_svg()
    create_docx()
    for path in (PNG_PATH, SVG_PATH, DOCX_PATH):
        print(path)


if __name__ == "__main__":
    main()

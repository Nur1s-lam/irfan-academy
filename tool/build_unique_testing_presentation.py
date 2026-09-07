from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "Отчет_участника_3_ГОСТ.docx"
LOGO = ROOT / "assets" / "images" / "irfan_logo.png"
OUTPUT = Path(
    os.environ.get(
        "IRFAN_PRESENTATION_OUTPUT",
        ROOT / "docs" / "Презентация_индивидуальной_работы_участник_3.pptx",
    )
)

WIDE = 13.333333
HIGH = 7.5
PX_W = 960
PX_H = 540

COLORS = {
    "ink": "18313F",
    "ink_2": "274A59",
    "paper": "FBF8F1",
    "cream": "F1E9D8",
    "sand": "D3A13B",
    "sand_light": "F4E7C5",
    "teal": "2F7375",
    "teal_light": "DCEAE6",
    "coral": "B95F4B",
    "coral_light": "F2DED8",
    "slate": "60727B",
    "line": "D7D0C3",
    "white": "FFFFFF",
    "black": "121A1D",
    "terminal": "111719",
    "success": "2E7D5B",
}

TITLE_FONT = "Georgia"
BODY_FONT = "Segoe UI"


def color(value: str) -> RGBColor:
    return RGBColor.from_string(COLORS.get(value, value).lstrip("#"))


def u(value: float):
    return Pt(value)


def rectangle(slide, x, y, w, h, fill, line=None, radius=False, transparency=0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, u(x), u(y), u(w), u(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.fill.transparency = transparency
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = u(1)
    else:
        shape.line.fill.background()
    return shape


def circle(slide, x, y, diameter, fill, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, u(x), u(y), u(diameter), u(diameter))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if line:
        shape.line.color.rgb = color(line)
    else:
        shape.line.fill.background()
    return shape


def line(slide, x1, y1, x2, y2, fill="line", width=1.25):
    shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, u(x1), u(y1), u(x2), u(y2))
    shape.line.color.rgb = color(fill)
    shape.line.width = u(width)
    return shape


def text(
    slide,
    value,
    x,
    y,
    w,
    h,
    size=18,
    fill="ink",
    bold=False,
    font=BODY_FONT,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    spacing=2,
    margin=0,
):
    box = slide.shapes.add_textbox(u(x), u(y), u(w), u(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = u(margin)
    frame.margin_right = u(margin)
    frame.margin_top = u(margin)
    frame.margin_bottom = u(margin)
    frame.vertical_anchor = valign
    for index, row in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = row
        paragraph.alignment = align
        paragraph.space_after = u(spacing)
        paragraph.font.name = font
        paragraph.font.size = u(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color(fill)
    return box


def label(slide, value, x, y, fill="teal"):
    width = max(92, len(value) * 6.8 + 26)
    rectangle(slide, x, y, width, 24, fill, radius=True)
    text(slide, value.upper(), x + 12, y + 5, width - 24, 14, 10, "white", True, spacing=0)
    return width


def base_slide(prs, number, section, background="paper", dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rectangle(slide, 0, 0, PX_W, PX_H, background)
    rectangle(slide, 0, 0, 13, PX_H, "sand")
    line(slide, 42, 500, 918, 500, "line", 0.75)
    text(slide, section.upper(), 42, 25, 460, 18, 10, "sand" if dark else "teal", True, spacing=0)
    text(slide, f"{number:02d}", 866, 20, 52, 31, 20, "sand", True, TITLE_FONT, PP_ALIGN.RIGHT, spacing=0)
    text(slide, "IRFAN ACADEMY  /  ПРОИЗВОДСТВЕННАЯ ПРАКТИКА", 42, 510, 520, 13, 8, "cream" if dark else "slate", True, spacing=0)
    return slide


def title(slide, value, x=42, y=57, w=790, h=61, size=27, fill="ink"):
    return text(slide, value, x, y, w, h, size, fill, True, TITLE_FONT, spacing=0)


def add_notes(slide, value: str):
    frame = slide.notes_slide.notes_text_frame
    frame.clear()
    paragraph = frame.paragraphs[0]
    paragraph.text = value.strip()
    paragraph.font.name = BODY_FONT
    paragraph.font.size = Pt(14)


def media_blob(name: str) -> bytes:
    with zipfile.ZipFile(REPORT) as archive:
        return archive.read(f"word/media/{name}")


def report_image(slide, name, x, y, w, h, background="white", line_color="line", pad=8):
    data = media_blob(name)
    with Image.open(BytesIO(data)) as source:
        source_w, source_h = source.size
    rectangle(slide, x, y, w, h, background, line_color)
    inner_w = w - pad * 2
    inner_h = h - pad * 2
    scale = min(inner_w / source_w, inner_h / source_h)
    image_w = source_w * scale
    image_h = source_h * scale
    image_x = x + (w - image_w) / 2
    image_y = y + (h - image_h) / 2
    return slide.shapes.add_picture(BytesIO(data), u(image_x), u(image_y), u(image_w), u(image_h))


def number_marker(slide, value, x, y, fill="sand", diameter=38):
    circle(slide, x, y, diameter, fill)
    text(slide, str(value), x, y + 8, diameter, 21, 14, "white", True, TITLE_FONT, PP_ALIGN.CENTER, spacing=0)


def check_row(slide, y, first, second, result="ПРОЙДЕН"):
    line(slide, 52, y + 46, 908, y + 46, "line", 0.75)
    circle(slide, 54, y + 8, 26, "teal")
    text(slide, "✓", 54, y + 11, 26, 15, 11, "white", True, BODY_FONT, PP_ALIGN.CENTER, spacing=0)
    text(slide, first, 96, y + 8, 260, 29, 15, "ink", True)
    text(slide, second, 365, y + 7, 390, 34, 14, "slate")
    text(slide, result, 785, y + 10, 115, 20, 10, "success", True, align=PP_ALIGN.RIGHT, spacing=0)


def build_presentation() -> Path:
    if not REPORT.exists():
        raise FileNotFoundError(REPORT)
    if not LOGO.exists():
        raise FileNotFoundError(LOGO)

    prs = Presentation()
    prs.slide_width = Inches(WIDE)
    prs.slide_height = Inches(HIGH)

    # 1. Титульный слайд
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rectangle(slide, 0, 0, PX_W, PX_H, "ink")
    rectangle(slide, 0, 0, 22, PX_H, "sand")
    circle(slide, 694, -122, 390, "ink_2")
    circle(slide, 755, 44, 198, "paper")
    slide.shapes.add_picture(str(LOGO), u(782), u(70), u(145), u(145))
    label(slide, "Производственная практика", 62, 54, "teal")
    text(slide, "ТЕСТИРОВАНИЕ", 62, 119, 615, 48, 30, "sand", True, BODY_FONT, spacing=0)
    text(slide, "видеоуроков и\nраздела «Намаз»", 60, 166, 625, 150, 37, "white", True, TITLE_FONT, spacing=1)
    line(slide, 63, 336, 545, 336, "sand", 2)
    text(slide, "Индивидуальная работа в проекте Irfan Academy", 62, 359, 590, 34, 18, "cream", False)
    text(slide, "Студент группы ПИ ан-2-23\n[ФИО третьего участника]", 62, 420, 420, 62, 16, "white")
    text(slide, "Бишкек  ·  2026", 740, 454, 178, 24, 14, "cream", True, align=PP_ALIGN.RIGHT)
    text(slide, "01", 874, 506, 44, 18, 11, "sand", True, TITLE_FONT, PP_ALIGN.RIGHT)
    add_notes(
        slide,
        "Здравствуйте. Я представляю свою часть командного проекта Irfan Academy. Я занимался проверкой загрузки и просмотра видеоуроков, а также раздела «Намаз». В этот раздел входят определение местоположения, время намазов, направление Кыблы и напоминания. Сначала я покажу, как проводил проверку, затем продемонстрирую основные экраны и в конце расскажу о полученных результатах.",
    )

    # 2. Роль
    slide = base_slide(prs, 2, "Моя работа в проекте")
    title(slide, "Что я проверял в приложении")
    text(slide, "Irfan Academy", 42, 143, 440, 42, 26, "teal", True, TITLE_FONT)
    text(
        slide,
        "Образовательное приложение на Flutter, в котором объединены учебные материалы, видеоуроки, задания, расписание занятий и религиозные функции.",
        42, 195, 420, 116, 18, "ink",
    )
    rectangle(slide, 518, 126, 390, 292, "cream")
    text(slide, "МОЯ ЧАСТЬ РАБОТЫ", 548, 153, 325, 20, 11, "coral", True)
    number_marker(slide, 1, 548, 195, "teal")
    text(slide, "Видеоуроки", 603, 198, 220, 25, 18, "ink", True, TITLE_FONT)
    text(slide, "Выбор, загрузка и просмотр видео", 603, 229, 255, 36, 15, "slate")
    number_marker(slide, 2, 548, 285, "sand")
    text(slide, "Религиозные функции", 603, 288, 260, 25, 18, "ink", True, TITLE_FONT)
    text(slide, "Местоположение, намазы, Кыбла и уведомления", 603, 319, 270, 42, 15, "slate")
    text(slide, "Главная задача — проверить обычную работу и поведение приложения при ошибках.", 42, 354, 420, 84, 19, "ink", True, TITLE_FONT)
    add_notes(
        slide,
        "Irfan Academy — это образовательное приложение, в котором есть учебные материалы, видеоуроки, задания и раздел с религиозными функциями. Я проверял не всё приложение, а две конкретные части. Первая — загрузка и просмотр видео. Вторая — раздел «Намаз»: местоположение, расписание, Кыбла и напоминания. Во время работы я смотрел, правильно ли выполняются действия пользователя и что происходит, если пропадает интернет, запрещён доступ к местоположению или на устройстве нет компаса.",
    )

    # 3. Цель
    slide = base_slide(prs, 3, "Задача")
    label(slide, "Цель", 42, 66, "coral")
    text(slide, "Проверить обычные действия\nи работу приложения при ошибках", 42, 112, 660, 98, 29, "ink", True, TITLE_FONT)
    line(slide, 42, 230, 918, 230, "line", 1)
    tasks = [
        ("01", "Обычная работа", "Повторить основные действия пользователя"),
        ("02", "Работа при ошибке", "Отключить интернет или запретить разрешение"),
        ("03", "Сложные случаи", "Проверить большие файлы и отсутствие компаса"),
        ("04", "Повторная проверка", "Ещё раз пройти всё после исправлений"),
    ]
    for index, (number, heading, body) in enumerate(tasks):
        x = 42 + (index % 2) * 438
        y = 261 + (index // 2) * 101
        text(slide, number, x, y, 45, 31, 18, "sand", True, TITLE_FONT)
        text(slide, heading, x + 58, y, 330, 26, 17, "ink", True, TITLE_FONT)
        text(slide, body, x + 58, y + 34, 332, 37, 14, "slate")
    add_notes(
        slide,
        "Моя цель была простой: проверить, что пользователь может выполнить нужное действие и понимает, что происходит на экране. Сначала я проходил обычный путь без ошибок. Затем специально создавал проблемы: отключал доступ к местоположению, проверял недоступную ссылку и учитывал отсутствие компаса. Отдельно проверял большие видео и переход к следующему дню после последнего намаза. Когда ошибка исправлялась, я снова повторял те же действия и проверял остальные основные функции.",
    )

    # 4. Объекты и среда
    slide = base_slide(prs, 4, "Подготовка")
    title(slide, "С чего я начал проверку")
    report_image(slide, "image1.png", 530, 105, 378, 350, "black", "ink", 6)
    text(slide, "ОСНОВНЫЕ ЧАСТИ ПРОЕКТА", 42, 137, 420, 20, 11, "coral", True)
    components = [
        "сервис загрузки видео",
        "данные о видеоуроке",
        "видеоплеер",
        "расчёт времени намазов",
        "определение Кыблы",
        "напоминания",
    ]
    for index, item in enumerate(components):
        y = 178 + index * 44
        text(slide, f"{index + 1:02d}", 42, y, 34, 21, 12, "sand", True, TITLE_FONT)
        text(slide, item, 86, y, 380, 25, 16, "ink", index < 3)
    text(slide, "Я нашёл файлы, которые отвечают за выбранные функции.", 530, 458, 378, 34, 12, "slate", False, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "Перед проверкой я посмотрел структуру проекта и нашёл файлы, которые отвечают за нужные функции. Для видео это загрузка файла, сохранение данных урока и плеер. Для раздела «Намаз» — расчёт времени, экран расписания, Кыбла и напоминания. Это помогло понять, откуда приложение получает каждое значение. Например, длительность видео должна определяться автоматически после загрузки, а не вводиться администратором вручную.",
    )

    # 5. Методика
    slide = base_slide(prs, 5, "Как проходила работа")
    title(slide, "Как я проверял каждую функцию")
    steps = [
        ("1", "Подготовил", "Открыл экран и выбрал данные"),
        ("2", "Выполнил", "Повторил действия пользователя"),
        ("3", "Сравнил", "Проверил ожидаемый результат"),
        ("4", "Записал", "Отметил итог или ошибку"),
        ("5", "Повторил", "Снова проверил после исправления"),
    ]
    line(slide, 93, 212, 865, 212, "sand", 3)
    for index, (number, heading, body) in enumerate(steps):
        x = 72 + index * 178
        circle(slide, x, 185, 54, "paper", "sand")
        text(slide, number, x, 199, 54, 22, 16, "ink", True, TITLE_FONT, PP_ALIGN.CENTER)
        text(slide, heading, x - 28, 261, 110, 25, 17, "ink", True, TITLE_FONT, PP_ALIGN.CENTER)
        text(slide, body, x - 40, 300, 134, 58, 14, "slate", False, align=PP_ALIGN.CENTER)
    rectangle(slide, 110, 389, 740, 76, "teal")
    text(slide, "Для каждой проверки я записывал, что должно произойти и что произошло на самом деле.", 138, 405, 684, 45, 16, "white", True, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "Каждую функцию я проверял по одинаковому порядку. Сначала запускал приложение и открывал нужный экран. Затем подготавливал данные: выбирал видео, включал или отключал интернет, разрешал или запрещал местоположение. После этого выполнял действия обычного пользователя. Я заранее записывал, какой результат должен быть, и сравнивал его с тем, что показало приложение. Если находил ошибку, записывал шаги. После исправления повторял их ещё раз, а затем быстро проверял соседние функции.",
    )

    # 6. Поток видео
    slide = base_slide(prs, 6, "Проверка видео")
    title(slide, "От выбора файла до просмотра урока")
    rectangle(slide, 42, 140, 876, 214, "ink")
    video_steps = [
        ("01", "Файл", "выбор на устройстве"),
        ("02", "Загрузка", "процент и повтор"),
        ("03", "Хранилище", "ссылка и данные"),
        ("04", "База", "сохранение урока"),
        ("05", "Плеер", "просмотр учеником"),
    ]
    for index, (number, heading, body) in enumerate(video_steps):
        x = 66 + index * 172
        text(slide, number, x, 166, 45, 29, 17, "sand", True, TITLE_FONT)
        text(slide, heading, x, 213, 145, 25, 17, "white", True, TITLE_FONT)
        text(slide, body, x, 254, 143, 43, 13, "cream")
        if index < 4:
            line(slide, x + 137, 229, x + 163, 229, "sand", 2)
    text(slide, "КОГДА ПРОВЕРКА ЗАВЕРШЕНА", 42, 392, 260, 21, 11, "coral", True)
    text(slide, "Урок появляется в списке и открывает тот же файл, который был выбран администратором.", 42, 424, 790, 46, 19, "ink", True, TITLE_FONT)
    add_notes(
        slide,
        "Проверка видео начиналась в админ-панели. Я выбирал файл с компьютера и смотрел, началась ли загрузка. Во время загрузки проверял процент и сообщение при ошибке. После завершения Cloudinary, где хранится видео, возвращал ссылку и данные о файле. В базе Firestore сохранялась только небольшая запись об уроке. Затем я входил как обычный пользователь, открывал список уроков и запускал загруженное видео. Проверка считалась успешной, если открывался именно выбранный файл.",
    )

    # 7. Выбор видео
    slide = base_slide(prs, 7, "Проверка видео · выбор файла")
    title(slide, "Выбор видео в админ-панели")
    report_image(slide, "image3.png", 493, 116, 415, 343, "white", "line", 8)
    text(slide, "Проверено", 42, 143, 230, 26, 13, "coral", True)
    points = [
        "открытие системного окна",
        "выбор MP4-файла",
        "отмена без создания урока",
        "автоматический запуск загрузки",
        "не нужно вводить ссылку вручную",
    ]
    for index, item in enumerate(points):
        y = 188 + index * 48
        circle(slide, 42, y + 1, 23, "teal")
        text(slide, "✓", 42, y + 4, 23, 13, 10, "white", True, align=PP_ALIGN.CENTER)
        text(slide, item, 80, y, 360, 27, 16, "ink")
    rectangle(slide, 42, 438, 394, 37, "sand_light")
    text(slide, "Результат: выбор файла работает", 56, 448, 365, 18, 12, "ink", True, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "Сначала я вошёл под учётной записью администратора и открыл раздел видеоуроков. Нажал кнопку выбора видео и проверил, что появилось обычное окно Windows. Затем выбрал MP4-файл. Загрузка началась сразу после выбора. Я также нажимал «Отмена»: в этом случае новый урок не появлялся. Ссылку и длительность вводить вручную не нужно. Администратор заполняет только номер, название и описание урока.",
    )

    # 8. Прогресс
    slide = base_slide(prs, 8, "Проверка видео · загрузка")
    title(slide, "Что происходит во время загрузки")
    report_image(slide, "image2.png", 42, 128, 544, 310, "black", "ink", 6)
    text(slide, "0—100%", 642, 145, 245, 56, 33, "teal", True, TITLE_FONT, PP_ALIGN.CENTER)
    text(slide, "непрерывное обновление", 642, 208, 245, 23, 13, "slate", False, align=PP_ALIGN.CENTER)
    line(slide, 631, 257, 897, 257, "line")
    text(slide, "до 4 попыток", 642, 282, 245, 40, 24, "sand", True, TITLE_FONT, PP_ALIGN.CENTER)
    text(slide, "если временно пропал интернет", 642, 332, 245, 42, 13, "slate", False, align=PP_ALIGN.CENTER)
    rectangle(slide, 626, 399, 276, 51, "teal_light")
    text(slide, "Окно загрузки не зависает навсегда", 645, 410, 238, 31, 13, "ink", True, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "После выбора файла я следил за окном загрузки. Процент должен постепенно меняться от нуля до ста, чтобы пользователь видел, что программа работает. Затем я временно отключал интернет и проверял сообщение об ошибке. Программа может попробовать отправить файл ещё раз, но делает это не больше четырёх раз. Если загрузка не удалась, окно не должно висеть бесконечно. После восстановления соединения загрузку можно запустить повторно.",
    )

    # 9. Размеры видео
    slide = base_slide(prs, 9, "Проверка видео · размер файла")
    title(slide, "Как загружаются небольшие и большие видео")
    text(slide, "95 МБ", 382, 134, 196, 74, 42, "sand", True, TITLE_FONT, PP_ALIGN.CENTER)
    line(slide, 105, 238, 855, 238, "ink", 3)
    circle(slide, 466, 226, 24, "sand")
    text(slide, "ДО 95 МБ", 72, 281, 330, 24, 14, "teal", True, align=PP_ALIGN.CENTER)
    text(slide, "Небольшое видео отправляется целиком. Во время передачи показывается процент загрузки.", 72, 326, 330, 75, 17, "ink", False, align=PP_ALIGN.CENTER)
    text(slide, "БОЛЕЕ 95 МБ", 558, 281, 330, 24, 14, "coral", True, align=PP_ALIGN.CENTER)
    text(slide, "Большое видео делится на части примерно по 8 МБ. Части отправляются по очереди.", 558, 326, 330, 75, 17, "ink", False, align=PP_ALIGN.CENTER)
    text(slide, "Большое видео обрабатывается по частям, поэтому приложение не зависает.", 160, 438, 640, 45, 17, "ink", True, TITLE_FONT, PP_ALIGN.CENTER)
    add_notes(
        slide,
        "Учебные видео могут быть длинными, поэтому я проверял файлы разного размера. Если файл меньше 95 мегабайт, он отправляется целиком. Если больше, программа делит его на части примерно по 8 мегабайт и отправляет их по очереди. Я проверял, что для каждого размера выбирается нужный способ, процент продолжает обновляться, а после последней части загрузка завершается. Такой подход не занимает всю память устройства одним большим файлом.",
    )

    # 10. Firestore
    slide = base_slide(prs, 10, "Проверка видео · сохранение")
    title(slide, "Какие данные сохраняются после загрузки")
    report_image(slide, "image4.png", 318, 126, 590, 300, "black", "ink", 5)
    text(slide, "В БАЗЕ ОСТАЮТСЯ", 42, 148, 230, 25, 15, "coral", True)
    fields = ["ссылка на видео", "номер файла в хранилище", "имя файла", "размер", "формат", "длительность"]
    for index, field in enumerate(fields):
        y = 195 + index * 39
        text(slide, "—", 42, y, 20, 22, 14, "sand", True)
        text(slide, field, 70, y, 210, 22, 15, "ink", True)
    text(slide, "Сохранённые данные совпадают с выбранным файлом", 318, 449, 590, 26, 14, "success", True, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "После загрузки я открыл базу Firestore и нашёл созданный урок. Затем сравнил сохранённые значения с исходным видео. В базе есть название урока, имя файла, размер, формат, длительность, ссылка и номер файла в хранилище. Само видео там не хранится — оно находится в Cloudinary. Благодаря этому база хранит только небольшую запись, а приложение получает видео по ссылке. Все значения на проверочном примере совпали с выбранным файлом.",
    )

    # 11. Плеер
    slide = base_slide(prs, 11, "Проверка видео · просмотр")
    title(slide, "Проверка видеоплеера")
    report_image(slide, "image5.png", 469, 111, 439, 352, "black", "ink", 5)
    controls = ["Запуск", "Пауза", "Перемотка", "Повтор", "Скорость", "Закрытие"]
    for index, item in enumerate(controls):
        x = 42 + (index % 2) * 188
        y = 153 + (index // 2) * 80
        text(slide, f"0{index + 1}", x, y, 34, 23, 13, "sand", True, TITLE_FONT)
        text(slide, item, x + 42, y, 137, 24, 16, "ink", True)
    rectangle(slide, 42, 401, 366, 57, "teal")
    text(slide, "После закрытия урока видео останавливается", 62, 414, 326, 33, 14, "white", True, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "Затем я вошёл в приложение как ученик, открыл список видеоуроков и выбрал загруженный урок. Проверил запуск и паузу, перемотал видео вперёд, включил повтор и изменил скорость. Также открыл неправильную ссылку, чтобы увидеть понятное сообщение об ошибке. После закрытия окна видео должно остановиться и не продолжать играть в фоне. Все основные кнопки сработали, а загруженный файл открылся полностью.",
    )

    # 12. Матрица видеомодуля
    slide = base_slide(prs, 12, "Проверка видео · ошибки")
    title(slide, "Что происходит, если возникла ошибка")
    text(slide, "ЧТО Я ДЕЛАЛ", 96, 140, 260, 17, 10, "coral", True)
    text(slide, "ЧТО ДОЛЖНО ПРОИЗОЙТИ", 365, 140, 390, 17, 10, "coral", True)
    text(slide, "СТАТУС", 785, 140, 115, 17, 10, "coral", True, align=PP_ALIGN.RIGHT)
    check_row(slide, 166, "Отмена выбора", "Урок не создаётся, форма остаётся доступной")
    check_row(slide, 222, "Временный сбой сети", "Сообщение об ошибке и кнопка «Повторить»")
    check_row(slide, 278, "Неправильная ссылка", "Плеер показывает ошибку и предлагает повторить")
    check_row(slide, 334, "Закрытие урока", "Видео останавливается и не играет в фоне")
    rectangle(slide, 52, 420, 856, 53, "sand_light")
    text(slide, "После исправлений я повторно прошёл эти четыре проверки.", 78, 434, 804, 24, 16, "ink", True, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "Здесь собраны проверки при ошибках. Сначала я отменил выбор файла и убедился, что пустой урок не появился. Затем временно отключил интернет: программа показала сообщение и не пыталась загружать бесконечно. После этого открыл урок с неправильной ссылкой — плеер показал ошибку и предложил повторить. В последней проверке я закрыл урок во время просмотра и убедился, что видео остановилось. После исправлений все четыре проверки были повторены.",
    )

    # 13. Религиозный модуль
    slide = base_slide(prs, 13, "Раздел «Намаз»")
    title(slide, "Как я проверял раздел «Намаз»")
    report_image(slide, "image6.png", 42, 120, 346, 353, "white", "line", 5)
    stages = [
        ("1", "Местоположение", "текущие координаты или Бишкек"),
        ("2", "Время", "расписание и направление Кыблы"),
        ("3", "Экран", "намазы, компас и обратный отсчёт"),
        ("4", "Напоминание", "уведомление на телефоне"),
    ]
    for index, (number, heading, body) in enumerate(stages):
        y = 137 + index * 82
        number_marker(slide, number, 443, y, "sand" if index % 2 else "teal", 33)
        text(slide, heading, 492, y, 185, 24, 16, "ink", True, TITLE_FONT)
        text(slide, body, 492, y + 30, 370, 34, 14, "slate")
    text(slide, "Каждый шаг проверялся отдельно, затем весь раздел целиком.", 443, 447, 445, 38, 14, "coral", True)
    add_notes(
        slide,
        "Раздел «Намаз» я проверял по порядку. Сначала смотрел, какое местоположение выбрало приложение. Затем проверял рассчитанное время намазов и направление Кыблы. После этого смотрел, правильно ли всё отображается на экране: список намазов, ближайший намаз, обратный отсчёт и компас. На телефоне дополнительно проверял напоминание. После отдельных проверок я ещё раз открыл весь раздел целиком и убедился, что данные согласованы между собой.",
    )

    # 14. Геолокация
    slide = base_slide(prs, 14, "Раздел «Намаз» · местоположение")
    title(slide, "Если местоположение недоступно")
    report_image(slide, "image7.png", 42, 143, 430, 118, "white", "line", 5)
    report_image(slide, "image8.png", 42, 304, 430, 135, "white", "line", 5)
    text(slide, "Проверенные состояния", 530, 143, 300, 25, 15, "coral", True)
    geo_cases = ["определение места включено", "доступ разрешён", "доступ запрещён", "доступ запрещён постоянно", "нет ответа в течение 5 секунд"]
    for index, item in enumerate(geo_cases):
        y = 188 + index * 43
        text(slide, f"{index + 1:02d}", 530, y, 31, 21, 12, "sand", True, TITLE_FONT)
        text(slide, item, 576, y, 280, 22, 16, "ink")
    rectangle(slide, 520, 420, 382, 44, "teal_light")
    text(slide, "Если место не определено: Бишкек · 42.8746, 74.5698", 538, 429, 346, 26, 13, "ink", True, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "Для проверки местоположения я сначала разрешил доступ и убедился, что приложение использует текущие координаты. Затем отключил службу и отдельно запретил разрешение. Также проверил случай, когда устройство не отвечает в течение пяти секунд. Во всех этих ситуациях экран не должен зависать. Если координаты получить не удалось, приложение использует данные Бишкека: 42.8746 и 74.5698. Название выбранного места показывается на экране, поэтому пользователь понимает, для какого города рассчитано время.",
    )

    # 15. Намазы
    slide = base_slide(prs, 15, "Раздел «Намаз» · расписание")
    title(slide, "Проверка времени и следующего намаза")
    rectangle(slide, 42, 145, 876, 66, "ink")
    text(slide, "Настройки расчёта", 70, 166, 255, 24, 15, "white", True, TITLE_FONT)
    text(slide, "Аср по ханафитскому методу", 326, 166, 282, 24, 14, "sand", True, TITLE_FONT, PP_ALIGN.CENTER)
    text(slide, "5 обязательных намазов", 609, 166, 275, 24, 15, "white", True, TITLE_FONT, PP_ALIGN.RIGHT)
    prayer_cases = [
        ("Обычный день", "Выбирается первый намаз, время которого ещё не наступило."),
        ("После Иша", "Следующим событием становится Фаджр следующего дня."),
        ("Смена даты", "Расписание и обратный отсчёт пересчитываются на экране."),
    ]
    for index, (heading, body) in enumerate(prayer_cases):
        x = 42 + index * 292
        text(slide, f"0{index + 1}", x, 258, 45, 28, 17, "sand", True, TITLE_FONT)
        text(slide, heading, x + 52, 258, 220, 25, 17, "ink", True, TITLE_FONT)
        line(slide, x, 301, x + 258, 301, "line")
        text(slide, body, x, 323, 258, 84, 15, "slate")
    text(slide, "После Иша следующим правильно показывается Фаджр следующего дня.", 104, 438, 752, 45, 16, "success", True, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "На экране расписания я проверял время пяти обязательных намазов, выделение ближайшего намаза и обратный отсчёт. Для расчёта используется метод Всемирной мусульманской лиги, а Аср определяется по ханафитскому методу. Самая важная отдельная проверка проводилась после Иша. В текущем дне больше нет будущих намазов, поэтому приложение должно показать Фаджр следующего дня. Название, дата и оставшееся время отобразились правильно.",
    )

    # 16. Кыбла
    slide = base_slide(prs, 16, "Раздел «Намаз» · Кыбла")
    title(slide, "Проверка направления Кыблы")
    report_image(slide, "image9.png", 42, 121, 473, 342, "white", "line", 5)
    text(slide, "245°", 617, 147, 190, 70, 43, "sand", True, TITLE_FONT, PP_ALIGN.CENTER)
    text(slide, "пример направления на Каабу", 602, 222, 220, 24, 13, "slate", False, align=PP_ALIGN.CENTER)
    line(slide, 566, 273, 858, 273, "line")
    text(slide, "При повороте устройства", 566, 302, 292, 25, 16, "ink", True, TITLE_FONT, PP_ALIGN.CENTER)
    text(slide, "я смотрел, правильно ли двигается стрелка и меняется угол", 579, 345, 266, 73, 15, "slate", False, align=PP_ALIGN.CENTER)
    text(slide, "Без компаса экран не закрывается с ошибкой", 568, 441, 288, 38, 13, "success", True, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "Для текущего места приложение рассчитывает угол направления на Каабу. Затем оно получает направление телефона от компаса. Я медленно поворачивал устройство и смотрел, правильно ли двигается стрелка. Отдельно проверил полный поворот, когда значение переходит от 359 градусов к нулю. Если компаса нет или ему нужна настройка, экран не должен закрываться с ошибкой. В этом случае пользователь видит понятное сообщение.",
    )

    # 17. Уведомления
    slide = base_slide(prs, 17, "Раздел «Намаз» · напоминания")
    title(slide, "Напоминания на телефоне и в браузере")
    text(slide, "ТЕЛЕФОН ANDROID", 42, 132, 400, 19, 11, "coral", True)
    report_image(slide, "image10.png", 42, 166, 410, 151, "white", "line", 5)
    text(slide, "Разрешение  ·  время  ·  название намаза", 42, 336, 410, 22, 13, "ink", True, align=PP_ALIGN.CENTER)
    text(slide, "БРАУЗЕР", 508, 132, 400, 19, 11, "coral", True)
    report_image(slide, "image11.png", 508, 166, 400, 221, "white", "line", 5)
    rectangle(slide, 508, 406, 400, 50, "teal")
    text(slide, "Браузер не запускает неподдерживаемое напоминание", 526, 416, 364, 32, 13, "white", True, align=PP_ALIGN.CENTER)
    text(slide, "На телефоне создаётся напоминание на выбранное время.", 42, 409, 410, 45, 14, "slate", False, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "Напоминания я проверял отдельно на телефоне и в браузере. На Android сначала разрешил уведомления, затем создал напоминание и сверил название намаза и время. Если телефон не разрешает точное срабатывание, используется обычный режим. В браузере такая функция не поддерживается. Поэтому программа не пытается её запускать и спокойно продолжает работу. На скриншоте видно, что раздел открывается без ошибки.",
    )

    # 18. Автоматизация
    slide = base_slide(prs, 18, "Последняя проверка")
    title(slide, "Финальная проверка проекта")
    report_image(slide, "image12.png", 42, 135, 430, 108, "terminal", "ink", 5)
    report_image(slide, "image13.png", 42, 279, 430, 108, "terminal", "ink", 5)
    metrics = [
        ("0", "ошибок в коде", "teal"),
        ("2 / 2", "автоматические проверки", "sand"),
        ("ГОТОВО", "проверка серверного файла", "coral"),
    ]
    for index, (value, caption, accent) in enumerate(metrics):
        x = 526 + (index % 2) * 186
        y = 140 + (index // 2) * 150
        text(slide, value, x, y, 158, 54, 31, accent, True, TITLE_FONT, PP_ALIGN.CENTER)
        text(slide, caption, x, y + 60, 158, 36, 13, "slate", False, align=PP_ALIGN.CENTER)
    rectangle(slide, 712, 290, 176, 92, "teal_light")
    text(slide, "ПРОЙДЕНО", 721, 307, 158, 31, 18, "success", True, TITLE_FONT, PP_ALIGN.CENTER)
    text(slide, "повторная проверка", 721, 347, 158, 20, 12, "slate", False, align=PP_ALIGN.CENTER)
    text(slide, "Команды проверили код, а основные функции я повторно проверил вручную.", 526, 424, 362, 47, 14, "ink", True, align=PP_ALIGN.CENTER)
    add_notes(
        slide,
        "В конце я запустил три команды. Первая команда проверила Flutter-код и не нашла ошибок. Вторая запустила две автоматические проверки: вход администратора и стартовый экран. Обе завершились успешно. Третья команда проверила серверный файл Firebase Functions и также не нашла ошибок. После команд я ещё раз вручную открыл загрузку видео, плеер, расписание намазов, Кыблу и напоминания, чтобы убедиться, что основные функции продолжают работать.",
    )

    # 19. Итог
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    rectangle(slide, 0, 0, PX_W, PX_H, "ink")
    rectangle(slide, 0, 0, 22, PX_H, "sand")
    slide.shapes.add_picture(str(LOGO), u(752), u(56), u(142), u(142))
    text(slide, "ИТОГ РАБОТЫ", 62, 55, 350, 21, 12, "sand", True)
    text(slide, "Проверка завершена:\nосновные функции работают", 60, 103, 625, 96, 30, "white", True, TITLE_FONT)
    results = [
        "видео выбирается, загружается и открывается",
        "данные о видео сохраняются отдельно от файла",
        "местоположение и время намазов определяются",
        "Кыбла работает даже без данных компаса",
        "напоминания учитывают телефон и браузер",
        "код и основные функции проверены повторно",
    ]
    for index, item in enumerate(results):
        col = index % 2
        row = index // 2
        x = 62 + col * 418
        y = 247 + row * 70
        number_marker(slide, index + 1, x, y, "teal" if col == 0 else "sand", 30)
        text(slide, item, x + 44, y + 2, 340, 42, 14, "cream", index in (0, 5))
    line(slide, 62, 468, 898, 468, "ink_2", 1)
    text(slide, "Спасибо за внимание", 62, 487, 350, 28, 17, "sand", True, TITLE_FONT)
    text(slide, "Готов ответить на вопросы комиссии", 548, 489, 350, 23, 13, "cream", False, align=PP_ALIGN.RIGHT)
    add_notes(
        slide,
        "В результате я проверил весь путь видеоурока: выбор файла, загрузку, сохранение данных и просмотр. Также проверил местоположение, время намазов, следующий намаз, Кыблу и напоминания. Ошибки сети, отказ разрешений и отсутствие компаса не должны закрывать приложение. Команды проверки кода и две автоматические проверки завершились успешно. На демонстрации я могу последовательно показать загрузку видео, работу плеера, расписание намазов, Кыблу и результаты команд. Спасибо за внимание.",
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_presentation())

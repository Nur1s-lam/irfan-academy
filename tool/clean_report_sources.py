import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "docs" / "Отчет_участника_3_ГОСТ.docx"

SOURCES = [
    "1. ГОСТ 7.32–2017. Отчёт о научно-исследовательской работе. Структура и правила оформления. — URL: https://protect.gost.ru/gost/details/7d280e43-7036-4a69-8e6e-d15867028343",
    "2. Кыргызский государственный технический университет им. И. Раззакова. Положение о порядке организации практик студентов. — URL: https://kstu.kg/fileadmin/user_upload/13._polozhenie_o_porjadke_organizacii_praktik_studentov_kgtu_im._i._razzakova__2020_g._.pdf",
    "3. Flutter. Официальная документация. — URL: https://docs.flutter.dev/",
    "4. Dart. Официальная документация языка и библиотек. — URL: https://dart.dev/guides",
    "5. Firebase Authentication for Flutter. Официальная документация. — URL: https://firebase.google.com/docs/auth/flutter/start",
    "6. Cloud Firestore for Flutter. Официальная документация. — URL: https://firebase.google.com/docs/firestore/quickstart",
    "7. Firebase Callable Functions. Официальная документация. — URL: https://firebase.google.com/docs/functions/callable",
    "8. Cloudinary Upload API Reference. — URL: https://cloudinary.com/documentation/image_upload_api_reference",
    "9. Cloudinary Chunked Asset Upload. — URL: https://cloudinary.com/documentation/upload_images#chunked_asset_upload",
    "10. Пакет video_player для Flutter. — URL: https://pub.dev/packages/video_player",
    "11. Пакет geolocator для Flutter. — URL: https://pub.dev/packages/geolocator",
    "12. Пакет adhan для Dart. — URL: https://pub.dev/packages/adhan",
    "13. Пакет flutter_compass для Flutter. — URL: https://pub.dev/packages/flutter_compass",
    "14. Пакет flutter_local_notifications для Flutter. — URL: https://pub.dev/packages/flutter_local_notifications",
]


def set_update_fields(document: Document) -> None:
    settings = document.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def replace_paragraph_text(paragraph, value: str) -> None:
    for run in paragraph.runs:
        run._element.getparent().remove(run._element)
    run = paragraph.add_run(value)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(14)


def format_source_paragraph(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fmt = paragraph.paragraph_format
    fmt.left_indent = Cm(1.25)
    fmt.first_line_indent = Cm(-0.75)
    fmt.right_indent = Cm(0)
    fmt.line_spacing = 1.5
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.keep_together = True
    fmt.widow_control = True
    for run in paragraph.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(14)
        run.font.bold = False


def clean_report() -> tuple[int, int]:
    document = Document(REPORT)

    heading_fixed = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text == "1.3 Назначение разрабатываемых модулей":
            replace_paragraph_text(paragraph, "1.1 Назначение разрабатываемых модулей")
            for run in paragraph.runs:
                run.font.bold = True
            heading_fixed += 1
            break

    source_heading_index = next(
        index
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip() == "СПИСОК ИСПОЛЬЗОВАННЫХ ИСТОЧНИКОВ"
    )
    source_paragraphs = []
    for paragraph in document.paragraphs[source_heading_index + 1 :]:
        if paragraph.style.name.startswith("Heading"):
            break
        if paragraph.text.strip():
            source_paragraphs.append(paragraph)

    if len(source_paragraphs) != len(SOURCES):
        raise RuntimeError(
            f"Ожидалось {len(SOURCES)} источников, найдено {len(source_paragraphs)}. "
            "Файл не изменён."
        )

    for paragraph, source in zip(source_paragraphs, SOURCES):
        clean_text = re.sub(r"\s+", " ", source).strip()
        replace_paragraph_text(paragraph, clean_text)
        format_source_paragraph(paragraph)

    heading = document.paragraphs[source_heading_index]
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    heading.paragraph_format.page_break_before = True
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(0)
    heading.paragraph_format.space_after = Pt(18)
    for run in heading.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        run.font.size = Pt(14)
        run.font.bold = True

    set_update_fields(document)
    document.save(REPORT)
    return heading_fixed, len(source_paragraphs)


if __name__ == "__main__":
    fixed, sources = clean_report()
    print(f"heading_fixed={fixed}")
    print(f"sources_formatted={sources}")
    print(REPORT)

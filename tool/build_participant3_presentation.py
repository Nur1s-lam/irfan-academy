from io import BytesIO
import os
from pathlib import Path
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = Path(os.environ.get(
    "IRFAN_PRESENTATION_OUTPUT",
    ROOT / "docs" / "Презентация_индивидуальной_работы_участник_3.pptx",
))
LOGO = ROOT / "assets" / "images" / "irfan_logo.png"
REPORT = ROOT / "docs" / "Отчет_участника_3_ГОСТ.docx"

COLORS = {
    "navy": "214F7B",
    "navy_dark": "173853",
    "blue": "3C78A8",
    "gold": "E2B92F",
    "cream": "FBF9F2",
    "white": "FFFFFF",
    "text": "243444",
    "muted": "667788",
    "border": "D9E1E8",
    "pale_blue": "EAF2F8",
    "pale_gold": "FCF5D7",
    "green": "2F855A",
    "pale_green": "E8F4ED",
    "red": "B94A48",
    "pale_red": "F9E9E8",
}


def rgb(name_or_hex: str) -> RGBColor:
    value = COLORS.get(name_or_hex, name_or_hex).lstrip("#")
    return RGBColor.from_string(value)


def pt(value: float):
    return Pt(value)


def add_shape(slide, shape_type, x, y, w, h, fill, line=None, line_width=1):
    shape = slide.shapes.add_shape(shape_type, pt(x), pt(y), pt(w), pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = rgb(line)
        shape.line.width = pt(line_width)
    return shape


def add_text(
    slide,
    value,
    x,
    y,
    w,
    h,
    size=18,
    color="text",
    bold=False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin=0,
    spacing=3,
):
    shape = slide.shapes.add_textbox(pt(x), pt(y), pt(w), pt(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = pt(margin)
    frame.margin_right = pt(margin)
    frame.margin_top = pt(margin)
    frame.margin_bottom = pt(margin)
    frame.vertical_anchor = valign
    for index, line in enumerate(value.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.alignment = align
        paragraph.space_after = pt(spacing)
        paragraph.font.name = "Times New Roman"
        paragraph.font.size = pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = rgb(color)
    return shape


def add_line(slide, x1, y1, x2, y2, color="blue", width=2):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, pt(x1), pt(y1), pt(x2), pt(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = pt(width)
    return line


def add_circle(slide, value, x, y, diameter, fill="navy", size=17, text_color="white"):
    circle = add_shape(slide, MSO_SHAPE.OVAL, x, y, diameter, diameter, fill)
    frame = circle.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.text = value
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.font.name = "Times New Roman"
    paragraph.font.size = pt(size)
    paragraph.font.bold = True
    paragraph.font.color.rgb = rgb(text_color)
    return circle


def add_chevron(slide, x, y, w=22, h=18, fill="gold"):
    return add_shape(slide, MSO_SHAPE.CHEVRON, x, y, w, h, fill)


def add_footer(slide, number):
    add_line(slide, 42, 509, 918, 509, "border", 0.75)
    add_text(slide, "IRFAN ACADEMY  |  ИНДИВИДУАЛЬНАЯ РАБОТА", 42, 514, 560, 15, 9, "muted", True)
    add_text(slide, str(number), 870, 511, 48, 17, 10, "navy", True, PP_ALIGN.RIGHT)


def add_base_slide(prs, title, section, number):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 960, 540, "cream")
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 16, 540, "gold")
    add_text(slide, section.upper(), 44, 20, 390, 21, 11, "blue", True)
    add_text(slide, title, 42, 43, 805, 44, 28, "navy_dark", True)
    add_line(slide, 42, 92, 170, 92, "gold", 4)
    slide.shapes.add_picture(str(LOGO), pt(870), pt(18), pt(50), pt(50))
    add_footer(slide, number)
    return slide


def add_card(slide, title, body, x, y, w, h, number=None, accent="blue", fill="white", body_size=17):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, "border")
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, 7, h, accent)
    title_x = x + 20
    if number is not None:
        add_circle(slide, str(number), x + 20, y + 17, 40, accent, 16)
        title_x = x + 72
    add_text(slide, title, title_x, y + 16, w - (title_x - x) - 18, 31, 20, "navy_dark", True)
    add_text(slide, body, x + 20, y + 58, w - 40, h - 70, body_size, "text")


def add_explanation(slide, value, x=42, y=420, w=876, h=67, size=17):
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, "pale_blue", "border")
    add_text(slide, value, x + 18, y + 5, w - 36, h - 10, size, "navy_dark", False, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE)


def add_screenshot(slide, number, label, x, y, w, h):
    frame = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, "white", "blue", 2)
    frame.fill.transparency = 0.08
    add_shape(slide, MSO_SHAPE.RECTANGLE, x + 14, y + 14, w - 28, 6, "gold")
    add_text(slide, "МЕСТО ДЛЯ СКРИНШОТА", x + 30, y + h / 2 - 34, w - 60, 25, 17, "blue", True, PP_ALIGN.CENTER)
    add_text(slide, label, x + 32, y + h / 2, w - 64, 48, 16, "muted", False, PP_ALIGN.CENTER)
    add_text(slide, f"Рисунок {number} — {label}", x, y + h + 8, w, 34, 13, "text", False, PP_ALIGN.CENTER)


def add_report_image(slide, media_name, caption, x, y, w, h, number=None):
    with zipfile.ZipFile(REPORT) as archive:
        data = archive.read(f"word/media/{media_name}")

    with Image.open(BytesIO(data)) as source:
        source_w, source_h = source.size

    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, "white", "border")
    inner_w = w - 14
    inner_h = h - 14
    scale = min(inner_w / source_w, inner_h / source_h)
    picture_w = source_w * scale
    picture_h = source_h * scale
    picture_x = x + (w - picture_w) / 2
    picture_y = y + (h - picture_h) / 2
    slide.shapes.add_picture(BytesIO(data), pt(picture_x), pt(picture_y), pt(picture_w), pt(picture_h))
    prefix = f"Рисунок {number} — " if number is not None else ""
    add_text(slide, f"{prefix}{caption}", x, y + h + 7, w, 29, 12, "muted", False, PP_ALIGN.CENTER)


def add_flow_step(slide, number, title, body, x, y, w=170):
    add_circle(slide, str(number), x + (w - 46) / 2, y, 46, "navy", 18)
    add_text(slide, title, x, y + 60, w, 29, 18, "navy_dark", True, PP_ALIGN.CENTER)
    add_text(slide, body, x, y + 96, w, 70, 15, "text", False, PP_ALIGN.CENTER)


def build_previous_presentation():
    if not LOGO.exists():
        raise FileNotFoundError(f"Логотип не найден: {LOGO}")

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    # 1. Титульный слайд
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 960, 540, "cream")
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 28, 540, "gold")
    add_shape(slide, MSO_SHAPE.OVAL, 700, -110, 350, 350, "pale_gold")
    add_shape(slide, MSO_SHAPE.OVAL, 778, 360, 250, 250, "pale_blue")
    slide.shapes.add_picture(str(LOGO), pt(708), pt(86), pt(196), pt(196))
    add_text(slide, "ПРОИЗВОДСТВЕННАЯ ПРАКТИКА", 64, 50, 470, 26, 13, "blue", True)
    add_text(slide, "Тестирование мультимедийных и религиозных модулей", 62, 100, 595, 145, 33, "navy_dark", True)
    add_text(slide, "Индивидуальная работа в проекте Irfan Academy", 65, 262, 555, 36, 20, "gold", True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 64, 326, 495, 1, "border")
    add_text(slide, "Студент группы ПИ ан-2-23\n[ФИО третьего участника]", 64, 347, 470, 64, 19, "text")
    add_text(slide, "Бишкек — 2026", 64, 464, 260, 27, 15, "muted", True)
    add_text(slide, "IRFAN ACADEMY", 700, 292, 212, 31, 20, "navy", True, PP_ALIGN.CENTER)
    add_footer(slide, 1)

    # 2. Цель и выполненные задачи
    slide = add_base_slide(prs, "Цель и состав тестирования", "01  Введение", 2)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 42, 123, 350, 275, "navy", "navy")
    add_text(slide, "Цель работы", 68, 149, 296, 32, 22, "gold", True, PP_ALIGN.CENTER)
    add_text(slide, "Моей целью было проверить корректность, устойчивость и удобство модулей Irfan Academy: видеоуроков, геолокации, расписания намазов, Кыблы и локальных уведомлений.", 70, 201, 292, 165, 19, "white")
    add_text(slide, "Выполненные задачи", 442, 124, 430, 31, 22, "navy_dark", True)
    add_text(slide, "• составлены положительные и негативные сценарии\n• проверены загрузка и воспроизведение видео\n• проверены геолокация, намазы и Кыбла\n• воспроизведены платформенные ошибки Web\n• проведена повторная проверка исправлений\n• выполнены автоматизированные проверки", 442, 174, 445, 240, 18, "text", spacing=8)
    add_explanation(slide, "В презентации показаны условия проверок, обнаруженные дефекты и фактические результаты повторного тестирования.", y=427, h=60, size=17)

    # 3. Зона ответственности
    slide = add_base_slide(prs, "Объекты и виды тестирования", "01  Введение", 3)
    add_card(slide, "Видеомодуль", "Функциональные, сетевые и граничные сценарии загрузки и плеера.", 42, 122, 420, 138, "01", "blue", body_size=17)
    add_card(slide, "Религиозный модуль", "Геолокация, расписание, следующий намаз и данные компаса.", 498, 122, 420, 138, "02", "gold", body_size=17)
    add_card(slide, "Платформенные проверки", "Android-уведомления и безопасная работа Web без zonedSchedule.", 42, 284, 420, 138, "03", "green", body_size=17)
    add_card(slide, "Регрессионный контроль", "Повтор дефектов, analyze, Flutter-тесты и node --check.", 498, 284, 420, 138, "04", "navy", body_size=17)
    add_explanation(slide, "Моя зона ответственности — проверка этих модулей и подтверждение результата после исправления найденных ошибок.", y=430, h=57, size=16)

    # 4. Технологии
    slide = add_base_slide(prs, "Среда и инструменты тестирования", "02  Методика", 4)
    technologies = [
        (42, 121, "Flutter / Dart", "Запуск приложения, analyze и автоматизированные тесты.", "blue"),
        (348, 121, "Firebase", "Проверка входа, ролей и метаданных в Firestore.", "gold"),
        (654, 121, "Cloudinary", "Проверка передачи видео и ответа сервиса.", "green"),
        (42, 278, "Firebase Functions", "Синтаксическая проверка серверного файла.", "green"),
        (348, 278, "GitHub", "Сопоставление исправлений и повторных проверок.", "blue"),
        (654, 278, "Android / Web", "Проверка различий геолокации, компаса и уведомлений.", "gold"),
    ]
    for x, y, title, body, accent in technologies:
        add_card(slide, title, body, x, y, 264, 124, accent=accent, body_size=16)
    add_explanation(slide, "Проверки выполнялись на действующей версии проекта. Для каждого сценария сравнивались ожидаемый и фактический результаты.", y=427, h=59, size=17)

    # 5. Проверяемый путь видеомодуля
    slide = add_base_slide(prs, "Модель тестирования видеомодуля", "03  Тестирование видео", 5)
    flow = [
        (42, "Действие", "выбор видео"),
        (220, "Клиент", "передача файла"),
        (398, "Cloudinary", "ответ и хранение"),
        (576, "Firestore", "проверка записи"),
        (754, "Плеер", "контроль просмотра"),
    ]
    for index, (x, title, body) in enumerate(flow):
        add_flow_step(slide, index + 1, title, body, x, 129, 155)
        if index < len(flow) - 1:
            add_line(slide, x + 145, 151, x + 168, 151, "gold", 2)
            add_chevron(slide, x + 162, 142, 17, 18)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 80, 317, 800, 110, "white", "border")
    add_text(slide, "Проверка охватывает полный путь данных: выбор файла, передачу, ответ Cloudinary, сохранение URL и метаданных в Firestore, открытие урока и работу сетевого плеера.", 105, 335, 750, 76, 18, "text")
    add_explanation(slide, "Отдельно проверено, что видео не передаётся через Firestore, а приложение не содержит секретного API-ключа Cloudinary.", y=438, h=50, size=16)

    # 6. Административная загрузка
    slide = add_base_slide(prs, "Проверка загрузки через административную панель", "03  Тестирование видео", 6)
    add_text(slide, "В положительном сценарии администратор выбирает видео с устройства. Проверяется автоматический запуск загрузки, отсутствие полей ручного ввода URL и длительности, а также доступность формы во время операции.", 42, 125, 390, 125, 19, "text")
    add_text(slide, "Также проверены отмена выбора, неподдерживаемый файл и повторное действие. После успешной загрузки технические данные должны заполняться автоматически.", 42, 274, 390, 106, 19, "text")
    add_screenshot(slide, 1, "Выбор видео в административной панели", 472, 120, 420, 296)

    # 7. Процесс загрузки
    slide = add_base_slide(prs, "Проверка прогресса и ошибок загрузки", "03  Тестирование видео", 7)
    add_screenshot(slide, 2, "Индикатор загрузки видеофайла", 42, 122, 430, 285)
    add_text(slide, "Во время передачи проверялось непрерывное обновление процента от 0 до 100 и понятное состояние интерфейса без зависания.", 516, 129, 390, 100, 19, "text")
    add_text(slide, "При временной сетевой ошибке контролировались сообщение пользователю и повторные попытки. Их количество ограничено четырьмя.", 516, 251, 390, 105, 19, "text")
    add_text(slide, "Критерий успеха: получены URL и метаданные, урок появился в списке и открылся повторно.", 516, 378, 390, 61, 18, "green", True)

    # 8. Большие файлы
    slide = add_base_slide(prs, "Проверка файлов разного размера", "03  Тестирование видео", 8)
    add_card(slide, "Файл до 95 MB", "Проверяется потоковый запрос, прогресс и отсутствие хранения полного файла в памяти.", 42, 124, 405, 151, accent="blue", fill="pale_blue", body_size=18)
    add_card(slide, "Файл более 95 MB", "Проверяется последовательная передача блоков по 8 MiB и завершение chunked upload.", 513, 124, 405, 151, accent="gold", fill="pale_gold", body_size=18)
    metrics = [(42, "8 MiB", "один блок"), (272, "4", "максимум попыток"), (502, "0–100%", "прогресс"), (732, "AUTO", "метаданные")]
    for x, value, label in metrics:
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 315, 186, 97, "white", "border")
        add_text(slide, value, x + 12, 329, 162, 38, 27, "navy", True, PP_ALIGN.CENTER)
        add_text(slide, label, x + 12, 375, 162, 24, 15, "muted", False, PP_ALIGN.CENTER)
    add_explanation(slide, "Граничные сценарии подтверждают выбор нужного режима по размеру файла и сохранение прогресса при передаче длинных видео.", y=428, h=59, size=17)

    # 9. Метаданные Firestore
    slide = add_base_slide(prs, "Проверка метаданных видеоурока", "03  Тестирование видео", 9)
    add_text(slide, "После загрузки сверялись поля:", 42, 126, 380, 29, 21, "navy_dark", True)
    add_text(slide, "• номер, название и описание урока\n• secure URL видео\n• public ID файла\n• размер и формат\n• длительность\n• дата создания", 50, 174, 360, 246, 19, "text", spacing=9)
    add_screenshot(slide, 3, "Документ видеоурока в Firestore", 455, 120, 437, 296)

    # 10. Плеер
    slide = add_base_slide(prs, "Проверка сетевого видеоплеера", "03  Тестирование видео", 10)
    add_screenshot(slide, 4, "Воспроизведение загруженного видеоурока", 42, 122, 480, 299)
    add_text(slide, "Исходный дефект: воспроизведение имитировалось таймером. После исправления повторная проверка подтвердила загрузку реального файла через VideoPlayerController.networkUrl.", 560, 126, 355, 114, 19, "text")
    add_text(slide, "Проверенные действия:", 560, 260, 330, 28, 20, "navy_dark", True)
    add_text(slide, "• запуск и паузу\n• перемотку и повтор\n• изменение скорости\n• отображение загрузки\n• сообщение об ошибке и повтор", 560, 300, 350, 125, 18, "text")

    # 11. Расчёт намазов
    slide = add_base_slide(prs, "Проверка геолокации и расчёта намазов", "04  Тестирование функций", 11)
    steps = [
        (42, "Координаты", "получение геолокации"),
        (270, "Контроль", "тайм-аут 5 секунд"),
        (498, "Расчёт", "MWL и ханафитский мазхаб"),
        (726, "Результат", "расписание на текущую дату"),
    ]
    for index, (x, title, body) in enumerate(steps):
        add_flow_step(slide, index + 1, title, body, x, 130, 190)
        if index < len(steps) - 1:
            add_line(slide, x + 175, 152, x + 205, 152, "gold", 2)
            add_chevron(slide, x + 198, 143, 18, 18)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 72, 315, 376, 112, "pale_gold", "border")
    add_text(slide, "Если координаты недоступны, используются резервные координаты Бишкека: 42.8746, 74.5698.", 94, 335, 332, 75, 18, "text", False, PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 512, 315, 376, 112, "pale_blue", "border")
    add_text(slide, "После Иша следующим событием определяется Фаджр следующего дня.", 534, 335, 332, 75, 18, "text", False, PP_ALIGN.CENTER)
    add_explanation(slide, "Проверены разрешённая и недоступная геолокация, резервные координаты, расчёт текущего дня и переход к следующему дню.", y=437, h=50, size=16)

    # 12. Экран намазов
    slide = add_base_slide(prs, "Проверка расписания и следующего намаза", "04  Тестирование функций", 12)
    add_text(slide, "Сверялось отображение пяти обязательных намазов, местоположения и выделение ближайшего события.", 42, 129, 385, 90, 20, "text")
    add_text(slide, "Проверены обратный отсчёт, обновление после смены даты и особый сценарий после Иша, когда следующим становится Фаджр следующего дня.", 42, 248, 385, 125, 19, "text")
    add_screenshot(slide, 5, "Экран расписания намазов", 466, 120, 426, 296)

    # 13. Кыбла
    slide = add_base_slide(prs, "Проверка направления Кыблы", "04  Тестирование функций", 13)
    add_screenshot(slide, 6, "Экран компаса Кыблы", 42, 120, 426, 296)
    add_text(slide, "Проверялось получение азимута Кыблы по текущим координатам и его сравнение с направлением датчика компаса.", 514, 129, 382, 113, 19, "text")
    add_text(slide, "При повороте устройства контролировались направление и плавное обновление стрелки.", 514, 270, 382, 86, 19, "text")
    add_text(slide, "Негативный сценарий: при отсутствии компаса показывается понятное сообщение без аварийного завершения.", 514, 370, 382, 75, 18, "green", True)

    # 14. Уведомления
    slide = add_base_slide(prs, "Проверка локальных уведомлений", "04  Тестирование функций", 14)
    add_card(slide, "Android", "Проверены разрешение, создание локального напоминания, точный режим и переход к обычному при ограничении платформы.", 42, 123, 405, 164, accent="green", fill="pale_green", body_size=18)
    add_card(slide, "Web", "Воспроизведена ошибка zonedSchedule. После исправления браузер пропускает неподдерживаемый вызов и продолжает работу.", 42, 311, 405, 164, accent="blue", fill="pale_blue", body_size=18)
    add_screenshot(slide, 7, "Локальное уведомление на Android", 493, 120, 399, 296)

    # 15. Обнаруженные и повторно проверенные ошибки
    slide = add_base_slide(prs, "Обнаруженные и повторно проверенные дефекты", "05  Результаты", 15)
    rows = [
        (122, "Имитация видео таймером", "Настоящий сетевой плеер"),
        (190, "Ручной ввод URL и длительности", "Автоматические метаданные"),
        (258, "Большой расход памяти", "Потоковая и блочная загрузка"),
        (326, "zonedSchedule падал в Web", "Проверка платформы до вызова"),
        (394, "После Иша не было события", "Фаджр следующего дня"),
    ]
    add_text(slide, "ОБНАРУЖЕННЫЙ ДЕФЕКТ", 63, 96, 320, 20, 12, "red", True, PP_ALIGN.CENTER)
    add_text(slide, "ПОВТОРНАЯ ПРОВЕРКА", 590, 96, 310, 20, 12, "green", True, PP_ALIGN.CENTER)
    for y, before, after in rows:
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 42, y, 360, 52, "pale_red", "border")
        add_text(slide, before, 57, y + 12, 330, 30, 16, "text", False, PP_ALIGN.CENTER)
        add_line(slide, 425, y + 26, 525, y + 26, "gold", 2.5)
        add_chevron(slide, 519, y + 17, 22, 18)
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 558, y, 360, 52, "pale_green", "border")
        add_text(slide, after, 573, y + 12, 330, 30, 16, "text", True, PP_ALIGN.CENTER)

    # 16. Автоматизированные проверки
    slide = add_base_slide(prs, "Результаты автоматизированных проверок", "06  Автоматизация", 16)
    checks = [
        (42, "0", "ошибок\nflutter analyze", "blue"),
        (235, "2 / 2", "автотеста\nпройдено", "green"),
        (428, "OK", "node --check\nFunctions", "gold"),
    ]
    for x, value, label, color in checks:
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 125, 166, 132, "white", "border")
        add_text(slide, value, x + 12, 140, 142, 43, 29, color, True, PP_ALIGN.CENTER)
        add_text(slide, label, x + 12, 191, 142, 51, 16, "text", False, PP_ALIGN.CENTER)
    add_text(slide, "Flutter analyze завершился без замечаний. Два существующих теста подтвердили распознавание административной роли и открытие стартового экрана. Отдельно проверен синтаксис Firebase Functions.", 42, 286, 552, 115, 19, "text")
    add_screenshot(slide, 8, "Результаты flutter analyze и flutter test", 630, 120, 262, 278)
    add_explanation(slide, "Проверка выполнена 07.09.2026 на актуальной версии проекта.", y=452, h=36, size=16)

    # 17. Демонстрация
    slide = add_base_slide(prs, "Порядок демонстрации тестирования", "07  Защита", 17)
    demo = [
        ("1", "Вход администратора", "Проверяю роль и доступ к панели."),
        ("2", "Выбор видео", "Проверяю автозапуск и прогресс."),
        ("3", "Сохранение урока", "Сверяю метаданные и список."),
        ("4", "Воспроизведение", "Проверяю паузу, перемотку и скорость."),
        ("5", "Намазы и Кыбла", "Проверяю расчёт и реакцию компаса."),
        ("6", "Проверка Web", "Подтверждаю отсутствие прежнего сбоя."),
    ]
    for index, (number, title, body) in enumerate(demo):
        row, col = divmod(index, 3)
        x, y = 42 + col * 296, 120 + row * 153
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 270, 126, "white", "border")
        add_circle(slide, number, x + 18, y + 17, 39, "gold" if index == 1 else "navy", 16)
        add_text(slide, title, x + 69, y + 16, 180, 32, 18, "navy_dark", True)
        add_text(slide, body, x + 18, y + 64, 234, 49, 16, "text")
    add_explanation(slide, "Во время защиты я называю условие теста, выполняю действие и показываю комиссии фактический результат.", y=432, h=56, size=16)

    # 18. Итоги
    slide = add_base_slide(prs, "Итоговые результаты тестирования", "08  Заключение", 18)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 42, 122, 560, 312, "white", "border")
    add_text(slide, "По результатам тестирования:", 70, 148, 500, 30, 21, "navy_dark", True)
    add_text(slide, "• подтверждены загрузка, прогресс и метаданные видео\n• подтверждена работа сетевого плеера\n• проверены намазы, геолокация и Кыбла\n• повторно проверены исправленные дефекты\n• Web работает без сбоя zonedSchedule\n• analyze, 2 теста и node --check успешны", 70, 196, 500, 218, 18, "text", spacing=7)
    slide.shapes.add_picture(str(LOGO), pt(661), pt(137), pt(196), pt(196))
    add_text(slide, "Irfan Academy", 637, 345, 244, 37, 27, "navy", True, PP_ALIGN.CENTER)
    add_text(slide, "Спасибо за внимание", 637, 397, 244, 31, 19, "gold", True, PP_ALIGN.CENTER)
    add_text(slide, "Готов ответить на вопросы комиссии", 565, 460, 353, 24, 15, "muted", False, PP_ALIGN.RIGHT)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


def build_previous_final_presentation():
    """Предыдущая редакция презентации, оставленная для истории оформления."""
    if not LOGO.exists():
        raise FileNotFoundError(f"Логотип не найден: {LOGO}")
    if not REPORT.exists():
        raise FileNotFoundError(f"Отчёт не найден: {REPORT}")

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    # 1. Титульный слайд
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 960, 540, "cream")
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, 28, 540, "gold")
    add_shape(slide, MSO_SHAPE.OVAL, 700, -110, 350, 350, "pale_gold")
    add_shape(slide, MSO_SHAPE.OVAL, 778, 360, 250, 250, "pale_blue")
    slide.shapes.add_picture(str(LOGO), pt(708), pt(86), pt(196), pt(196))
    add_text(slide, "ПРОИЗВОДСТВЕННАЯ ПРАКТИКА", 64, 50, 470, 26, 13, "blue", True)
    add_text(slide, "Тестирование мультимедийных и религиозных модулей", 62, 100, 595, 145, 33, "navy_dark", True)
    add_text(slide, "Индивидуальная работа в проекте Irfan Academy", 65, 262, 555, 36, 20, "gold", True)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 64, 326, 495, 1, "border")
    add_text(slide, "Студент группы ПИ ан-2-23\n[ФИО третьего участника]", 64, 347, 470, 64, 19, "text")
    add_text(slide, "Бишкек — 2026", 64, 464, 260, 27, 15, "muted", True)
    add_text(slide, "IRFAN ACADEMY", 700, 292, 212, 31, 20, "navy", True, PP_ALIGN.CENTER)
    add_footer(slide, 1)

    # 2. Проект и роль участника
    slide = add_base_slide(prs, "Проект и моя роль", "01  Введение", 2)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 42, 123, 876, 135, "white", "border")
    add_text(
        slide,
        "Irfan Academy — образовательное приложение на Flutter. Оно объединяет учебные материалы, видеоуроки, расписание занятий и религиозные функции. В командном проекте моей основной задачей было тестирование тех частей, которые зависят от сети, внешних сервисов, геолокации и возможностей устройства.",
        66, 145, 828, 93, 18, "text",
    )
    add_card(slide, "Мультимедийный модуль", "Проверка выбора, загрузки, метаданных и воспроизведения видеоуроков.", 42, 286, 420, 134, "01", "blue", body_size=17)
    add_card(slide, "Религиозный модуль", "Проверка геолокации, намазов, Кыблы и локальных уведомлений.", 498, 286, 420, 134, "02", "gold", body_size=17)
    add_explanation(slide, "На защите я показываю не только готовые экраны, но и объясняю, какие условия проверялись и какой результат был получен.", y=435, h=52, size=16)

    # 3. Цель и задачи
    slide = add_base_slide(prs, "Цель и задачи тестирования", "01  Введение", 3)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 42, 122, 350, 298, "navy", "navy")
    add_text(slide, "Цель", 70, 148, 294, 31, 23, "gold", True, PP_ALIGN.CENTER)
    add_text(
        slide,
        "Проверить корректность, устойчивость и удобство мультимедийных и религиозных функций на Android и Web, а затем подтвердить отсутствие повторных ошибок после исправлений.",
        70, 194, 294, 202, 20, "white", False, PP_ALIGN.LEFT, MSO_ANCHOR.MIDDLE,
    )
    add_text(slide, "Основные задачи", 438, 124, 440, 31, 22, "navy_dark", True)
    add_text(
        slide,
        "• определить позитивные, негативные и граничные сценарии\n• проверить полный путь видео от выбора до просмотра\n• проверить работу функций с разрешениями и датчиками\n• воспроизвести платформенные ошибки Android и Web\n• повторно проверить исправленные дефекты\n• выполнить статический анализ и автоматизированные тесты",
        438, 171, 442, 248, 18, "text", spacing=8,
    )
    add_explanation(slide, "Для каждого сценария фиксировались начальные условия, действия, ожидаемый результат и фактическое поведение приложения.", y=438, h=49, size=16)

    # 4. Объекты и среда
    slide = add_base_slide(prs, "Объекты и среда тестирования", "02  Подготовка", 4)
    add_text(
        slide,
        "Перед проверкой я изучил структуру проекта и выделил компоненты, влияющие на нужные сценарии: сервис загрузки, модель видеоурока, сетевой плеер, PrayerService, виджет Кыблы и NotificationService.",
        42, 125, 400, 110, 18, "text",
    )
    add_text(slide, "Средства проверки", 42, 256, 380, 29, 20, "navy_dark", True)
    add_text(
        slide,
        "• Flutter и Dart — запуск и анализ клиента\n• Firebase — роли и метаданные Firestore\n• Cloudinary — передача и хранение видео\n• Android и Web — платформенные сценарии\n• терминал — analyze, test и node --check",
        48, 300, 388, 154, 17, "text", spacing=7,
    )
    add_report_image(slide, "image1.png", "Структура тестируемых компонентов проекта", 488, 116, 380, 330, 1)

    # 5. Методика
    slide = add_base_slide(prs, "Как проводилось тестирование", "02  Методика", 5)
    flow = [
        (42, "Подготовка", "условия и данные"),
        (220, "Действие", "шаги пользователя"),
        (398, "Ожидание", "критерий успеха"),
        (576, "Проверка", "фактический результат"),
        (754, "Регрессия", "повтор после исправления"),
    ]
    for index, (x, title, body) in enumerate(flow):
        add_flow_step(slide, index + 1, title, body, x, 124, 155)
        if index < len(flow) - 1:
            add_line(slide, x + 145, 147, x + 168, 147, "gold", 2)
            add_chevron(slide, x + 162, 138, 17, 18)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 72, 327, 816, 112, "white", "border")
    add_text(
        slide,
        "Функциональные сценарии проверяли обычную работу пользователя. Негативные сценарии охватывали отказ разрешений, отсутствие датчика и сетевые ошибки. Граничные проверки использовались для больших файлов и перехода к следующему дню. После исправления выполнялась повторная проверка исходных шагов.",
        98, 339, 764, 88, 18, "text",
    )
    add_explanation(slide, "Такой порядок позволил получить воспроизводимые результаты, а не ограничиваться визуальным просмотром экранов.", y=445, h=50, size=16)

    # 6. Полный путь видео
    slide = add_base_slide(prs, "Полный сценарий проверки видеоурока", "03  Видеомодуль", 6)
    video_flow = [
        (42, "Выбор", "файл с устройства"),
        (220, "Загрузка", "прогресс и повторы"),
        (398, "Cloudinary", "URL и метаданные"),
        (576, "Firestore", "сохранение записи"),
        (754, "Плеер", "реальное видео"),
    ]
    for index, (x, title, body) in enumerate(video_flow):
        add_flow_step(slide, index + 1, title, body, x, 126, 155)
        if index < len(video_flow) - 1:
            add_line(slide, x + 145, 149, x + 168, 149, "gold", 2)
            add_chevron(slide, x + 162, 140, 17, 18)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 72, 327, 816, 112, "pale_blue", "border")
    add_text(
        slide,
        "Проверка начиналась в административной панели и завершалась открытием урока учеником. Я контролировал, что сам видеофайл передаётся в Cloudinary, а Firestore получает только описание и технические метаданные. После сохранения тот же URL должен использоваться сетевым плеером.",
        98, 347, 764, 74, 18, "navy_dark",
    )
    add_explanation(slide, "Критерий успеха — опубликованный урок появляется в списке и воспроизводится без ручного ввода ссылки или длительности.", y=445, h=50, size=16)

    # 7. Выбор файла
    slide = add_base_slide(prs, "Проверка выбора видеофайла", "03  Видеомодуль", 7)
    add_text(
        slide,
        "В административной панели я проверил открытие системного окна, выбор MP4-файла и отмену операции. После подтверждения выбора загрузка должна запускаться автоматически. Поля для ручного ввода ссылки и длительности отсутствуют, поэтому администратор вводит только учебные сведения: номер, название и описание урока.",
        42, 126, 390, 188, 18, "text",
    )
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 42, 327, 390, 112, "pale_green", "border")
    add_text(slide, "Результат: системный выбор файла открывается, отмена не создаёт урок, а подтверждённый файл автоматически переходит к загрузке.", 62, 340, 350, 88, 17, "green", True)
    add_report_image(slide, "image3.png", "Выбор файла с устройства", 470, 118, 422, 317, 2)

    # 8. Прогресс и сеть
    slide = add_base_slide(prs, "Прогресс и сетевые ошибки", "03  Видеомодуль", 8)
    add_report_image(slide, "image2.png", "Индикатор загрузки видео", 42, 121, 466, 294, 3)
    add_text(
        slide,
        "Во время передачи контролировалось обновление процента от 0 до 100. Индикатор показывает, что приложение продолжает работу и не зависло. При временной сетевой ошибке пользователь получает понятное сообщение, а сервис выполняет ограниченные повторные попытки — не более четырёх.",
        548, 127, 350, 190, 18, "text",
    )
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 548, 325, 350, 127, "pale_green", "border")
    add_text(slide, "Фактический результат: прогресс отображается, завершение определяется корректно, а ошибка не оставляет форму в бесконечном состоянии загрузки.", 568, 338, 310, 104, 17, "green", True)

    # 9. Большие файлы
    slide = add_base_slide(prs, "Проверка видео разных размеров", "03  Видеомодуль", 9)
    add_card(slide, "До 95 MB", "Файл передаётся потоковым запросом. Проверялись прогресс, завершение и использование памяти.", 42, 122, 405, 150, "01", "blue", "pale_blue", 17)
    add_card(slide, "Более 95 MB", "Файл делится на блоки по 8 MiB. Проверялись порядок частей и завершение chunked upload.", 513, 122, 405, 150, "02", "gold", "pale_gold", 17)
    metrics = [(42, "8 MiB", "размер блока"), (272, "4", "повторные попытки"), (502, "0–100%", "индикатор"), (732, "AUTO", "метаданные")]
    for x, value, label in metrics:
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, 312, 186, 91, "white", "border")
        add_text(slide, value, x + 12, 326, 162, 34, 25, "navy", True, PP_ALIGN.CENTER)
        add_text(slide, label, x + 12, 368, 162, 23, 14, "muted", False, PP_ALIGN.CENTER)
    add_explanation(slide, "Граничная проверка подтвердила выбор нужного режима по размеру файла. Видео не передаётся через Firestore и не загружается целиком в оперативную память.", y=428, h=59, size=17)

    # 10. Метаданные
    slide = add_base_slide(prs, "Проверка метаданных в Firestore", "03  Видеомодуль", 10)
    add_text(
        slide,
        "После успешной загрузки я открыл документ videoLessons и сверил данные с выбранным файлом. В записи присутствуют название, описание, имя файла, размер, тип содержимого, длительность, secure URL и идентификатор Cloudinary. Сам бинарный файл в базе отсутствует.",
        42, 126, 374, 174, 18, "text",
    )
    add_text(slide, "Проверяемые поля", 42, 320, 360, 28, 20, "navy_dark", True)
    add_text(slide, "videoUrl  •  videoPath  •  fileName\nfileSize  •  contentType  •  duration", 48, 363, 355, 66, 17, "blue", True, PP_ALIGN.CENTER)
    add_report_image(slide, "image4.png", "Документ видеоурока в Firestore", 448, 118, 444, 310, 4)

    # 11. Плеер
    slide = add_base_slide(prs, "Проверка сетевого видеоплеера", "03  Видеомодуль", 11)
    add_report_image(slide, "image5.png", "Воспроизведение опубликованного видеоурока", 42, 116, 424, 337, 5)
    add_text(
        slide,
        "Урок открывался из пользовательского списка по сохранённому адресу Cloudinary. Я последовательно проверил запуск, паузу, перемотку, повтор и изменение скорости. Отдельно контролировались индикатор буферизации, сообщение при недоступном ресурсе и повторная попытка открытия.",
        506, 126, 390, 185, 18, "text",
    )
    add_text(slide, "Проверенные действия", 506, 329, 370, 27, 20, "navy_dark", True)
    add_text(slide, "Запуск  •  Пауза  •  Перемотка\nПовтор  •  Скорость  •  Закрытие", 510, 372, 370, 58, 17, "blue", True, PP_ALIGN.CENTER)
    add_explanation(slide, "Результат: воспроизводится реальный сетевой файл, а контроллер освобождает ресурсы после закрытия урока.", y=459, h=29, size=15)

    # 12. Негативные сценарии видеомодуля
    slide = add_base_slide(prs, "Негативные сценарии видеомодуля", "03  Видеомодуль", 12)
    add_text(slide, "Негативные проверки показывают, как приложение ведёт себя при ошибке пользователя или внешнего сервиса. Во всех случаях интерфейс должен сохранить управляемое состояние.", 42, 111, 876, 55, 17, "text")
    headers = ["Сценарий", "Ожидаемая реакция", "Результат"]
    widths = [250, 465, 135]
    xs = [42, 292, 757]
    for x, width, header in zip(xs, widths, headers):
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, 181, width, 42, "navy", "white")
        add_text(slide, header, x + 8, 191, width - 16, 22, 15, "white", True, PP_ALIGN.CENTER)
    test_rows = [
        ("Отмена выбора файла", "Урок не создаётся, форма остаётся доступной", "Пройден"),
        ("Временный сбой сети", "Показ ошибки и ограниченный повтор запроса", "Пройден"),
        ("Недоступный URL", "Сообщение плеера и возможность повторить", "Пройден"),
        ("Закрытие плеера", "Остановка и освобождение контроллера", "Пройден"),
    ]
    for row_index, row in enumerate(test_rows):
        y = 223 + row_index * 55
        fill = "white" if row_index % 2 == 0 else "pale_blue"
        for x, width, value in zip(xs, widths, row):
            add_shape(slide, MSO_SHAPE.RECTANGLE, x, y, width, 55, fill, "border")
            add_text(slide, value, x + 9, y + 7, width - 18, 40, 15, "green" if x == xs[-1] else "text", x == xs[-1], PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_explanation(slide, "Критические сценарии повторялись после исправлений и затем включались в общую регрессионную проверку.", y=454, h=34, size=15)

    # 13. Религиозный модуль
    slide = add_base_slide(prs, "Состав проверки религиозного модуля", "04  Религиозные функции", 13)
    add_report_image(slide, "image6.png", "Экран намазов и направления Кыблы", 42, 115, 386, 342, 6)
    add_text(
        slide,
        "Проверка начиналась с получения координат. На их основе сервис рассчитывает расписание и азимут Кыблы. Затем интерфейс показывает пять обязательных намазов, ближайшее событие и обратный отсчёт. На Android дополнительно проверялись локальные напоминания, а на Web — безопасный пропуск неподдерживаемого метода.",
        474, 126, 420, 195, 18, "text",
    )
    add_text(slide, "Основные контрольные точки", 474, 343, 400, 27, 20, "navy_dark", True)
    add_text(slide, "Координаты  →  Расписание  →  Следующий намаз\nКыбла  →  Компас  →  Уведомление", 478, 387, 400, 63, 17, "blue", True, PP_ALIGN.CENTER)

    # 14. Геолокация
    slide = add_base_slide(prs, "Проверка геолокации", "04  Религиозные функции", 14)
    add_text(
        slide,
        "Я проверил включённую и отключённую службу геолокации, разрешённый и запрещённый доступ, окончательный запрет и превышение пятисекундного тайм-аута. При доступных координатах расписание рассчитывается для текущего местоположения. При ошибке используются резервные координаты Бишкека — 42.8746 и 74.5698.",
        42, 124, 398, 198, 18, "text",
    )
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 42, 342, 398, 96, "pale_green", "border")
    add_text(slide, "Результат: отказ разрешения или тайм-аут не приводит к сбою, а выбранное местоположение явно отображается на экране.", 62, 354, 358, 68, 17, "green", True)
    add_report_image(slide, "image7.png", "Отображение местоположения", 480, 121, 412, 128, 7)
    add_report_image(slide, "image8.png", "Карточка следующего намаза", 480, 290, 412, 146, 8)

    # 15. Расписание намазов
    slide = add_base_slide(prs, "Проверка расписания и следующего намаза", "04  Религиозные функции", 15)
    add_text(
        slide,
        "Расписание рассчитывается по методу Muslim World League, а время Асра — по ханафитскому мазхабу. На экране я сверил пять обязательных намазов, формат времени, выделение ближайшего события и обновление обратного отсчёта.",
        42, 122, 876, 88, 18, "text",
    )
    scenarios = [
        (42, "Обычный день", "Следующим выбирается первый намаз, время которого ещё не наступило.", "blue"),
        (338, "После Иша", "Следующим событием становится Фаджр следующего дня.", "gold"),
        (634, "Смена даты", "Расписание и обратный отсчёт пересчитываются без перезапуска экрана.", "green"),
    ]
    for index, (x, title, body, accent) in enumerate(scenarios, 1):
        add_card(slide, title, body, x, 244, 270, 158, str(index), accent, "white", 17)
    add_explanation(slide, "Граничный сценарий после Иша подтверждён: дата следующего события меняется, а обратный отсчёт остаётся положительным.", y=431, h=56, size=17)

    # 16. Кыбла
    slide = add_base_slide(prs, "Проверка направления Кыблы", "04  Религиозные функции", 16)
    add_report_image(slide, "image9.png", "Направление Кыблы по данным компаса", 42, 117, 432, 321, 9)
    add_text(
        slide,
        "Для текущих координат вычисляется абсолютный азимут на Каабу. Затем он сравнивается с направлением устройства. При повороте телефона я контролировал изменение стрелки и значения около границы 0/360 градусов, где особенно важна правильная нормализация угла.",
        516, 126, 380, 187, 18, "text",
    )
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 516, 338, 380, 100, "pale_green", "border")
    add_text(slide, "Если компас отсутствует или ожидает калибровки, экран показывает понятное состояние и не завершается с ошибкой.", 538, 355, 336, 66, 17, "green", True)

    # 17. Уведомления и Web
    slide = add_base_slide(prs, "Проверка уведомлений на Android и Web", "04  Религиозные функции", 17)
    add_report_image(slide, "image10.png", "Локальные уведомления приложения", 42, 118, 408, 135, 10)
    add_report_image(slide, "image11.png", "Работа Web-версии без платформенного сбоя", 42, 291, 500, 146, 11)
    add_text(
        slide,
        "На Android проверялись разрешение, создание канала, время срабатывания и название намаза. Сначала применяется точный режим, а при платформенном ограничении — обычный. Прошедшие события повторно не планируются.",
        580, 126, 316, 145, 17, "text",
    )
    add_text(
        slide,
        "В браузере метод zonedSchedule не поддерживается. Проверка kIsWeb выполняется до вызова, поэтому Web-версия продолжает работу без необработанного исключения.",
        580, 306, 316, 114, 17, "green", True,
    )

    # 18. Автоматизированные проверки
    slide = add_base_slide(prs, "Автоматизированные и регрессионные проверки", "05  Контроль качества", 18)
    add_report_image(slide, "image12.png", "flutter analyze: замечаний не обнаружено", 42, 118, 410, 104, 12)
    add_report_image(slide, "image13.png", "flutter test: два теста выполнены успешно", 42, 266, 410, 104, 13)
    checks = [
        (500, 120, "0", "ошибок flutter analyze", "blue"),
        (704, 120, "2 / 2", "теста Flutter", "green"),
        (500, 281, "OK", "node --check Functions", "gold"),
        (704, 281, "PASS", "ручная регрессия", "navy"),
    ]
    for x, y, value, label, color in checks:
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, 174, 116, "white", "border")
        add_text(slide, value, x + 14, y + 15, 146, 38, 25, color, True, PP_ALIGN.CENTER)
        add_text(slide, label, x + 14, y + 65, 146, 36, 15, "text", False, PP_ALIGN.CENTER)
    add_explanation(slide, "Автотесты проверяют административную роль и стартовый экран. Сетевые и аппаратные сценарии дополнительно проверены вручную в фактической среде.", y=430, h=57, size=16)

    # 19. Итог
    slide = add_base_slide(prs, "Итоговые результаты", "06  Заключение", 19)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 42, 120, 590, 326, "white", "border")
    add_text(slide, "По результатам моей работы подтверждено:", 68, 145, 535, 30, 21, "navy_dark", True)
    add_text(
        slide,
        "• видео выбирается с устройства и загружается с прогрессом\n• большие файлы передаются потоково или блоками\n• метаданные сохраняются отдельно от видеофайла\n• опубликованный урок воспроизводится сетевым плеером\n• геолокация, намазы и Кыбла обрабатывают граничные состояния\n• Android-уведомления работают, а Web не вызывает zonedSchedule\n• статический анализ, тесты и повторная регрессия завершены успешно",
        70, 193, 528, 229, 17, "text", spacing=6,
    )
    slide.shapes.add_picture(str(LOGO), pt(690), pt(143), pt(165), pt(165))
    add_text(slide, "Irfan Academy", 652, 329, 242, 34, 26, "navy", True, PP_ALIGN.CENTER)
    add_text(slide, "Спасибо за внимание", 652, 380, 242, 28, 19, "gold", True, PP_ALIGN.CENTER)
    add_text(slide, "На демонстрации я последовательно покажу загрузку видео, работу плеера, расписание намазов, Кыблу и результаты проверок.", 650, 414, 246, 78, 15, "muted", False, PP_ALIGN.CENTER)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    return OUTPUT


def build_presentation():
    from build_unique_testing_presentation import build_presentation as build_unique

    return build_unique()


if __name__ == "__main__":
    print(build_presentation())
